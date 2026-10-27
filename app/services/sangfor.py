"""Batch-based Sangfor registration using the existing order and bank ledger.

Order.extra_data owns the typed, versioned application. No second payment or
refund state machine is introduced; financial decisions remain in OfflinePayment.
"""
from copy import deepcopy
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
from zoneinfo import ZoneInfo
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import and_, case, func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import aliased

from app.adapter.database import get_db_ctx
from app.domain.order.src.index import Order, apply_order_status_transition
from app.domain.plan.src.index import Plan
from app.models.cert_product import CertProduct
from app.models.offline_payment import OfflineBankEntry, OfflinePayment, OfflinePaymentAudit, OfflinePaymentSubmission
from app.port.config import settings
from app.port.exceptions import BusinessException, ConflictException, NotFoundException
from app.schemas.sangfor import SangforCreate, SangforDecision, SangforResubmit, SangforSubmission
from app.services.agreement_template import ensure_accepted
from app.services.certification_identity import lock_verified_identity
from app.services.offline_payment import OfflinePaymentService
from app.services.plan_enrollment import PlanEnrollmentService
from app.services.registration_payment_guard import payment_filter, payment_summary, require_review_payment
from app.services.sangfor_catalog import DIRECTIONS, HEADERS, require_direction
from app.utils.payment import generate_out_trade_no


def now():
    return datetime.now(timezone.utc)


def is_sangfor(order):
    return (order.extra_data or {}).get("business_type") == "sangfor"


def registration_status(order):
    if order.status == "closed":
        return "cancelled"
    if order.status == "refunded":
        return "refunded_closed"
    if order.status == "pending":
        return "pending_payment"
    return (order.extra_data or {}).get("registration_state", "pending_review")


def registration_status_expression():
    refund = select(OfflinePaymentAudit.id).where(OfflinePaymentAudit.order_id == Order.id,
        OfflinePaymentAudit.action == "admin.generic_refund_approved").exists()
    return case((Order.status == "closed", "cancelled"), (Order.status == "refunded", "refunded_closed"),
        (Order.status == "pending", "pending_payment"),
        (and_(Order.status.in_(("paid", "completed")), refund), "refund_processing"),
        else_=func.coalesce(Order.extra_data["registration_state"].as_string(), "pending_review"))


def export_workbook(snapshots):
    from pathlib import Path
    template = Path(__file__).resolve().parents[1] / "assets" / "sangfor_registration.xlsx"
    workbook = load_workbook(template)
    sheet = workbook["Sheet1"]
    if [sheet.cell(2, column).value for column in range(1, 14)] != HEADERS:
        raise BusinessException("深信服报送模板表头不匹配，请联系管理员")
    # The source contains one yellow example row; replace it with actual records.
    for row in sheet.iter_rows(min_row=3, max_row=max(sheet.max_row, 3), max_col=13):
        for cell in row:
            cell.value = None
    for index, data in enumerate(snapshots, 3):
        date = datetime.fromisoformat(data["exam_date"]).astimezone(ZoneInfo(settings.APP_TIMEZONE)).replace(tzinfo=None)
        values = [data["name"], data["id_card"], data["phone"], date, data["mailing_address"], data["email"],
                  data["organization"], data["first_name"], data["last_name"], data["exam_direction"],
                  data["exam_voucher_code"], data["remark1"], data["remark2"]]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(index, column, value)
            if column == 4:
                cell.number_format = "yyyy-mm-dd"
            elif isinstance(value, str):
                cell.data_type = "s"  # User text must never become an Excel formula.
                cell.number_format = "@"
    # Keep the supplied layout; remove conflicting, incomplete dropdown rules.
    sheet.data_validations.dataValidation = []
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class SangforService:
    async def require_product(self, code):
        require_direction(code)
        async with get_db_ctx() as db:
            product = await db.scalar(select(CertProduct).where(CertProduct.code == code, CertProduct.type == "sangfor"))
            if product is None:
                raise NotFoundException("深信服认证产品")

    @staticmethod
    def query():
        return select(Order).where(Order.order_kind == "certification", Order.extra_data["business_type"].as_string() == "sangfor")

    async def batch(self, db, batch_id, *, public=False):
        pair = (await db.execute(select(Plan, CertProduct).join(CertProduct, CertProduct.code == Plan.product_type)
                                 .where(Plan.id == batch_id, CertProduct.type == "sangfor"))).first()
        if pair is None or (public and (not pair[1].is_active or pair[0].status not in {"published", "registration_closed"} or pair[1].code not in DIRECTIONS)):
            raise NotFoundException("深信服考试批次")
        return pair

    async def batch_statistics(self, db, ids, *, public=False):
        occupied = dict((await db.execute(select(Order.plan_id, func.count()).where(Order.plan_id.in_(ids),
            Order.status.in_(("pending", "paid", "completed"))).group_by(Order.plan_id))).all())
        counts = {}
        if not public and ids:
            statuses = self.query().with_only_columns(Order.plan_id, registration_status_expression().label("state")).where(Order.plan_id.in_(ids)).subquery()
            for plan_id, state, count in (await db.execute(select(statuses.c.plan_id, statuses.c.state, func.count()).group_by(statuses.c.plan_id, statuses.c.state))).all():
                counts.setdefault(plan_id, {})[state] = count
        return occupied, counts

    @staticmethod
    def batch_response(plan, product, *, enrolled=0, counts=None):
        result = {"id": plan.id, "product_code": product.code, "product_name": product.chinese_name,
                "description": plan.description or product.description, "name": plan.name, "status": plan.status,
                "exam_direction": DIRECTIONS.get(product.code, product.code), "apply_start": plan.apply_start,
                "apply_end": plan.apply_end, "exam_date": plan.exam_date, "exam_location": plan.exam_location,
                "price_cents": plan.price_cents, "capacity": plan.capacity, "enrolled": enrolled,
                "remaining_count": max(0, plan.capacity - enrolled) if plan.capacity else -1}
        if counts is not None:
            result["registration_counts"] = counts
        return result

    async def batches(self, *, public=False, product_code=None):
        async with get_db_ctx() as db:
            query = select(Plan, CertProduct).join(CertProduct, CertProduct.code == Plan.product_type).where(CertProduct.type == "sangfor")
            if public:
                query = query.where(CertProduct.is_active.is_(True), CertProduct.code.in_(DIRECTIONS), Plan.status.in_(("published", "registration_closed")))
            if product_code:
                query = query.where(Plan.product_type == product_code)
            pairs = (await db.execute(query.order_by(Plan.id.desc()))).all()
            occupied, counts = await self.batch_statistics(db, [plan.id for plan, _ in pairs], public=public)
            return [self.batch_response(plan, product, enrolled=occupied.get(plan.id, 0), counts=None if public else counts.get(plan.id, {})) for plan, product in pairs]

    async def get_batch(self, batch_id, *, public=False):
        async with get_db_ctx() as db:
            plan, product = await self.batch(db, batch_id, public=public)
            occupied, counts = await self.batch_statistics(db, [batch_id], public=public)
            return self.batch_response(plan, product, enrolled=occupied.get(batch_id, 0), counts=None if public else counts.get(batch_id, {}))

    async def refund_batches(self):
        """Minimal batch selector for ledger readers, without content permissions."""
        async with get_db_ctx() as db:
            rows = (await db.execute(select(Plan.id, Plan.name).join(CertProduct,
                CertProduct.code == Plan.product_type).where(CertProduct.type == "sangfor").order_by(Plan.id.desc()))).mappings().all()
            return [dict(row) for row in rows]

    @staticmethod
    async def refund_approved(db, order_id):
        return bool(await db.scalar(select(OfflinePaymentAudit.id).where(OfflinePaymentAudit.order_id == order_id,
                              OfflinePaymentAudit.action == "admin.generic_refund_approved").limit(1)))

    async def response(self, db, order, *, brief=False):
        extra = order.extra_data
        payment = await payment_summary(db, order)
        status = registration_status(order)
        if order.status in {"paid", "completed"} and await self.refund_approved(db, order.id):
            status = "refund_processing"
            payment.review_payment_allowed = False
            payment.review_block_reason = "退款已审批，等待登记实际退款，报名已冻结"
        return {"id": order.id, "order_id": order.id, "batch_id": order.plan_id, "registration_no": order.out_trade_no,
                "status": status, "order_status": order.status, "price_cents": order.price,
                "revision": extra["revision"], "current_version": extra["current_version"],
                "candidate_snapshot": extra["candidate_snapshot"], "batch_snapshot": extra["batch_snapshot"],
                "versions": [] if brief else extra["versions"], "reviews": extra["reviews"][-1:] if brief else extra["reviews"], "exports": extra["exports"],
                "correction_due_at": extra.get("correction_due_at"), "payment_summary": payment.model_dump(),
                "created_at": order.created_at, "updated_at": order.updated_at}

    async def lock_order(self, db, order_id, user_id=None):
        query = self.query().where(Order.id == order_id)
        if user_id is not None:
            query = query.where(Order.user_id == user_id)
        order = await db.scalar(query.with_for_update())
        if order is None:
            raise NotFoundException("深信服报名")
        return order

    async def create(self, user_id, data: SangforCreate):
        try:
            async with get_db_ctx() as db:
                async with db.begin():
                    plan = await PlanEnrollmentService().lock_enrollable_plan(db, plan_id=data.batch_id, user_id=user_id, expected_vendor="sangfor")
                    direction = require_direction(plan.product_type)
                    if not plan.exam_date:
                        raise BusinessException("该批次未配置考试时间")
                    identity = await lock_verified_identity(db, user_id=user_id)
                    await ensure_accepted(db, user_id, "cert_registration", message="请先阅读并同意认证报名信息处理授权协议")
                    snapshot = data.model_dump(exclude={"batch_id"}) | {"name": identity.real_name, "id_card": identity.id_card_number,
                        "exam_direction": direction, "exam_direction_code": plan.product_type, "exam_date": plan.exam_date.isoformat()}
                    order = Order(user_id=user_id, order_kind="certification", product_type=plan.product_type, plan_id=plan.id,
                        candidate_name=identity.real_name, candidate_idcard=identity.id_card_number, candidate_phone=data.phone,
                        price=plan.price_cents, status="pending", out_trade_no=generate_out_trade_no("SF"),
                        extra_data={"business_type": "sangfor", "revision": 1, "current_version": 1,
                            "registration_state": "pending_review", "candidate_snapshot": snapshot,
                            "batch_snapshot": {"name": plan.name, "product_code": plan.product_type,
                                "exam_direction": direction, "exam_date": plan.exam_date.isoformat(), "exam_location": plan.exam_location},
                            "versions": [{"version": 1, "submitted_at": now().isoformat(), "snapshot": snapshot}], "reviews": [], "exports": []})
                    db.add(order)
                    await db.flush()
                    await OfflinePaymentService.initialize(db, order, plan)
                    await db.flush()
                    await db.refresh(order)
                    return await self.response(db, order)
        except IntegrityError as error:
            raise ConflictException("已存在该批次的有效报名，请在我的报名查看") from error

    async def detail(self, order_id, user_id=None):
        async with get_db_ctx() as db:
            return await self.response(db, await self.lock_order(db, order_id, user_id))

    async def registrations(self, *, user_id=None, batch_id=None, state=None, keyword=None, payment_state=None, refund_only=False, page=1, page_size=20):
        async with get_db_ctx() as db:
            query = self.query()
            if user_id is not None:
                query = query.where(Order.user_id == user_id)
            if batch_id:
                await self.batch(db, batch_id)
                query = query.where(Order.plan_id == batch_id)
            if keyword:
                query = query.where(Order.candidate_name.ilike(f"%{keyword}%") | Order.candidate_phone.ilike(f"%{keyword}%") | Order.out_trade_no.ilike(f"%{keyword}%"))
            if payment_state:
                query = query.where(payment_filter(payment_state))
            if state:
                query = query.where(registration_status_expression() == state)
            if refund_only:
                query = query.where(registration_status_expression().in_(("refund_processing", "refunded_closed")))
            total = await db.scalar(select(func.count()).select_from(query.subquery()))
            rows = list((await db.scalars(query.order_by(Order.id.desc()).offset((page - 1) * page_size).limit(page_size))).all())
            return {"items": [await self.response(db, row, brief=True) for row in rows], "total": total, "page": page, "page_size": page_size}

    async def refunds(self, *, batch_id=None, registration_id=None, keyword=None, request_kind=None,
                      requested_at_start=None, requested_at_end=None, group=None, status=None, page=1, page_size=20):
        """Project the existing bank ledger without inventing refund requests.

        Ordinary paid orders enter this list only after refund approval. Unallocated
        transfers use their own exception path, including completed returns whose
        order is closed rather than refunded. All joins remain one row per order.
        """
        if page < 1 or not 1 <= page_size <= 100:
            raise BusinessException("分页参数不正确，页码须大于零，每页最多100条")
        if any(value is not None and value <= 0 for value in (batch_id, registration_id)):
            raise BusinessException("请选择有效的批次或报名记录")
        for value in (requested_at_start, requested_at_end):
            if value is not None and (value.tzinfo is None or value.utcoffset() is None):
                raise BusinessException("退款日期筛选须包含时区")
        if requested_at_start and requested_at_end and requested_at_start > requested_at_end:
            raise BusinessException("开始时间不能晚于结束时间")
        if group not in {None, "pending", "processing", "abnormal"} or status not in {None, "approved", "exception", "succeeded"}:
            raise BusinessException("退款状态筛选不正确")
        if request_kind not in {None, "normal", "exception"}:
            raise BusinessException("退款类型筛选不正确")

        audit = OfflinePaymentAudit
        is_approval = audit.action == "admin.generic_refund_approved"
        is_mismatch = and_(audit.action == "admin.confirm", audit.summary["result"].as_string() == "exception")
        is_exception = or_(audit.action.in_(("user.exception", "admin.exception")), is_mismatch)
        events = select(
            audit.order_id.label("order_id"),
            func.min(case((is_approval, audit.created_at))).label("approved_at"),
            func.max(case((is_approval, audit.id))).label("approval_id"),
            func.min(case((is_exception, audit.created_at))).label("exception_at"),
            func.max(case((is_exception, audit.id))).label("exception_id"),
            func.max(case((is_mismatch, audit.id))).label("mismatch_id"),
            func.max(case((audit.action == "admin.exception_refund", audit.id))).label("resolution_id"),
        ).where(audit.order_id.is_not(None)).group_by(audit.order_id).subquery("sangfor_refund_events")
        approval, exception, mismatch, resolution = (aliased(audit) for _ in range(4))
        receipt, refund = aliased(OfflineBankEntry), aliased(OfflineBankEntry)
        active_exception = and_(Order.status.in_(("pending", "closed")), Order.offline_status == "exception", refund.id.is_(None))
        exception_kind = or_(active_exception, events.c.resolution_id.is_not(None))
        state = case((refund.id.is_not(None), "succeeded"), (active_exception, "exception"), else_="approved")
        requested_at = case((exception_kind, func.coalesce(events.c.exception_at, refund.created_at,
            Order.offline_submitted_at, Order.created_at)), else_=func.coalesce(events.c.approved_at, refund.created_at, Order.created_at))
        # A mismatched confirmation records an actual verified amount in its audit,
        # whereas the current submission contains only the user's declared amount.
        verified_exception_amount = case((mismatch.summary["version"].as_integer() == OfflinePayment.current_version,
                                          mismatch.summary["amount_cents"].as_integer()))
        amount = func.coalesce(refund.amount_cents, receipt.amount_cents,
            case((exception_kind, func.coalesce(verified_exception_amount,
                OfflinePaymentSubmission.data["amount_cents"].as_integer())), else_=Order.price))
        amount_source = case((or_(refund.amount_cents.is_not(None), receipt.amount_cents.is_not(None)), "ledger"),
            (and_(exception_kind, verified_exception_amount.is_not(None)), "verified"),
            (and_(exception_kind, OfflinePaymentSubmission.data["amount_cents"].as_integer().is_not(None)), "submission"),
            (~exception_kind, "order"), else_="unknown")
        base = select(
            Order.id.label("id"), Order.id.label("registration_id"), Order.id.label("order_id"),
            Order.plan_id.label("batch_id"), Plan.name.label("batch_name"),
            func.coalesce(CertProduct.chinese_name, Order.product_type).label("product_label"),
            Order.payment_channel.label("payment_channel"), Order.out_trade_no.label("out_trade_no"),
            Order.out_trade_no.label("registration_no"), Order.candidate_name.label("candidate_name"),
            case((exception_kind, "exception"), else_="normal").label("request_kind"), state.label("status"),
            amount.label("amount_cents"), amount_source.label("amount_source"),
            case((exception_kind, func.coalesce(resolution.summary["note"].as_string(),
                exception.summary["note"].as_string(), OfflinePayment.message)),
                else_=approval.summary["note"].as_string()).label("reason_detail"),
            requested_at.label("requested_at"),
            case((~exception_kind, events.c.approved_at)).label("approved_at"),
            refund.created_at.label("succeeded_at"),
        ).select_from(Order).join(OfflinePayment, OfflinePayment.order_id == Order.id).join(Plan, Plan.id == Order.plan_id
        ).outerjoin(CertProduct, CertProduct.code == Order.product_type
        ).outerjoin(events, events.c.order_id == Order.id
        ).outerjoin(approval, approval.id == events.c.approval_id
        ).outerjoin(exception, exception.id == events.c.exception_id
        ).outerjoin(mismatch, mismatch.id == events.c.mismatch_id
        ).outerjoin(resolution, resolution.id == events.c.resolution_id
        ).outerjoin(receipt, and_(receipt.order_id == Order.id, receipt.kind == "receipt")
        ).outerjoin(refund, and_(refund.order_id == Order.id, refund.kind == "refund")
        ).outerjoin(OfflinePaymentSubmission, and_(OfflinePaymentSubmission.order_id == Order.id,
            OfflinePaymentSubmission.version == OfflinePayment.current_version)
        ).where(Order.order_kind == "certification", Order.extra_data["business_type"].as_string() == "sangfor",
            Order.payment_channel == "bank_transfer", or_(active_exception,
                and_(events.c.approval_id.is_not(None), Order.status.in_(("paid", "completed"))),
                and_(refund.id.is_not(None), or_(Order.status == "refunded", events.c.resolution_id.is_not(None)))))
        if batch_id is not None:
            base = base.where(Order.plan_id == batch_id)
        if registration_id is not None:
            base = base.where(Order.id == registration_id)
        if keyword and keyword.strip():
            pattern = f"%{keyword.strip()}%"
            base = base.where(or_(Order.candidate_name.ilike(pattern), Order.candidate_phone.ilike(pattern),
                Order.out_trade_no.ilike(pattern)))
        projection = base.subquery("sangfor_refunds")
        filtered = select(projection)
        if request_kind:
            filtered = filtered.where(projection.c.request_kind == request_kind)
        if requested_at_start is not None:
            filtered = filtered.where(projection.c.requested_at >= requested_at_start)
        if requested_at_end is not None:
            filtered = filtered.where(projection.c.requested_at <= requested_at_end)
        stats = filtered.subquery()
        async with get_db_ctx() as db:
            counts = (await db.execute(select(stats.c.status, func.count()).group_by(stats.c.status))).all()
            if status:
                filtered = filtered.where(projection.c.status == status)
            elif group == "processing":
                filtered = filtered.where(projection.c.status == "approved")
            elif group == "abnormal":
                filtered = filtered.where(projection.c.status == "exception")
            elif group == "pending":
                filtered = filtered.where(False)
            total = await db.scalar(select(func.count()).select_from(filtered.subquery())) or 0
            rows = (await db.execute(filtered.order_by(projection.c.requested_at.asc(), projection.c.id.asc())
                .offset((page - 1) * page_size).limit(page_size))).mappings().all()
            return {"items": [dict(row) | {"last_error": None} for row in rows], "total": int(total),
                "page": page, "page_size": page_size, "status_counts": {key: int(value) for key, value in counts}}

    async def editable(self, db, order):
        await require_review_payment(db, order)
        if await self.refund_approved(db, order.id):
            raise ConflictException("退款已审批，不能继续报名或审核")
        plan, _product = await self.batch(db, order.plan_id)
        if plan.status in {"cancelled", "archived", "finalized"}:
            raise ConflictException("批次已终止，不能继续办理")

    @staticmethod
    def check_revision(order, expected):
        if order.extra_data["revision"] != expected:
            raise ConflictException("报名记录已更新，请刷新后重新操作")

    async def resubmit(self, order_id, user_id, data: SangforResubmit):
        async with get_db_ctx() as db:
            async with db.begin():
                order = await self.lock_order(db, order_id, user_id)
                self.check_revision(order, data.expected_revision)
                await self.editable(db, order)
                extra = deepcopy(order.extra_data)
                if registration_status(order) != "returned" or not extra.get("correction_due_at") or datetime.fromisoformat(extra["correction_due_at"]) < now():
                    raise ConflictException("当前不在有效补正期限内")
                identity = await lock_verified_identity(db, user_id=user_id)
                if identity.real_name != order.candidate_name or identity.id_card_number != order.candidate_idcard:
                    raise ConflictException("本人实名信息已变更，请联系管理员核实，不能替换报名人")
                snapshot = extra["candidate_snapshot"] | data.model_dump(exclude={"expected_revision"})
                extra["current_version"] += 1
                extra["revision"] += 1
                extra["registration_state"] = "pending_review"
                extra["candidate_snapshot"] = snapshot
                extra["correction_due_at"] = None
                extra["exported_version"] = None
                extra["versions"].append({"version": extra["current_version"], "submitted_at": now().isoformat(), "snapshot": snapshot})
                order.extra_data = extra
                order.candidate_phone = data.phone
                await db.flush()
                await db.refresh(order)
                return await self.response(db, order)

    async def review(self, order_id, admin_id, data: SangforDecision):
        async with get_db_ctx() as db:
            async with db.begin():
                order = await self.lock_order(db, order_id)
                self.check_revision(order, data.expected_revision)
                await self.editable(db, order)
                extra = deepcopy(order.extra_data)
                state = registration_status(order)
                if data.action in {"approve", "return"} and state != "pending_review":
                    raise ConflictException("仅待初审报名可执行该操作")
                if data.action.startswith("final_") and (state != "approved" or extra.get("exported_version") != extra["current_version"]):
                    raise ConflictException("当前版本须初审通过并已成功导出后才能登记外审结果")
                if data.action in {"return", "final_return"}:
                    if not data.reason or not data.correction_due_at or data.correction_due_at.tzinfo is None or data.correction_due_at <= now():
                        raise BusinessException("退回补正必须填写原因和未来的补正截止时间")
                    extra["registration_state"] = "returned"
                    extra["correction_due_at"] = data.correction_due_at.isoformat()
                else:
                    extra["registration_state"] = "approved" if data.action == "approve" else "final_approved"
                    if data.action == "final_approve":
                        apply_order_status_transition(order, "completed")
                extra["revision"] += 1
                extra["reviews"].append({"action": data.action, "reason": data.reason, "admin_id": admin_id,
                    "version": extra["current_version"], "created_at": now().isoformat(), "correction_due_at": extra.get("correction_due_at")})
                order.extra_data = extra
                await db.flush()
                await db.refresh(order)
                return await self.response(db, order)

    async def export_records(self, batch_id):
        async with get_db_ctx() as db:
            await self.batch(db, batch_id)
            rows = list(await db.scalars(self.query().where(Order.plan_id == batch_id).order_by(Order.id)))
            records = {}
            for order in rows:
                for entry in order.extra_data.get("exports", []):
                    key = entry.get("export_id") or entry["sha256"]
                    record = records.setdefault(key, {"id": key, "sha256": entry["sha256"], "created_at": entry["created_at"], "people": []})
                    record["created_at"] = min(record["created_at"], entry["created_at"])
                    record["people"].append({"registration_id": order.id, "registration_no": order.out_trade_no,
                        "name": order.candidate_name, "version": entry["version"], "is_current": entry["version"] == order.extra_data["current_version"]})
            return [record | {"people_count": len(record["people"])} for record in sorted(records.values(), key=lambda item: item["created_at"], reverse=True)]

    async def export(self, batch_id, ids, admin_id, expected_versions=None):
        async with get_db_ctx() as db:
            async with db.begin():
                await self.batch(db, batch_id)
                orders = list((await db.scalars(self.query().where(Order.plan_id == batch_id, Order.id.in_(ids)).order_by(Order.id).with_for_update())).all())
                if len(orders) != len(ids):
                    raise ConflictException("选择中包含不存在或属于其他批次的报名")
                snapshots = []
                for order in orders:
                    if expected_versions is not None and expected_versions.get(order.id) != order.extra_data["current_version"]:
                        raise ConflictException("所选报名资料版本已变化，请刷新名单后重新选择")
                    await self.editable(db, order)
                    if registration_status(order) not in {"approved", "final_approved"}:
                        raise ConflictException(f"{order.candidate_name}尚未初审通过，不能报送")
                    snapshot = order.extra_data["candidate_snapshot"]
                    SangforSubmission.model_validate({key: snapshot[key] for key in SangforSubmission.model_fields})
                    require_direction(snapshot["exam_direction_code"])
                    snapshots.append(snapshot)
                content = export_workbook(snapshots)
                checksum = sha256(content).hexdigest()
                export_id = uuid4().hex
                exported_at = now().isoformat()
                for order in orders:
                    extra = deepcopy(order.extra_data)
                    extra["exported_version"] = extra["current_version"]
                    extra["revision"] += 1
                    extra["exports"].append({"export_id": export_id, "version": extra["current_version"], "sha256": checksum, "admin_id": admin_id, "created_at": exported_at})
                    order.extra_data = extra
                return content
