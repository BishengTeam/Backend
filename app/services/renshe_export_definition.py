"""Current reporting contract. Product editors cannot change these rules.

Material stems follow the client's established filenames. New degree batches
require a one-inch white-background portrait; historical batches keep their rules.
Extensions always describe the actual file.
"""

import hashlib
import inspect
import json
import re
from datetime import date
from pathlib import Path

from app.port.exceptions import BusinessException

REGISTRATION_TEMPLATE_NAME = "报名信息.xlsx"
WORK_HISTORY_TEMPLATE_NAME = "工作经历.xlsx"
MATERIAL_STEMS = {
    "work_proof": "工龄证明",
    "id_card_front": "证件正面",
    "id_card_back": "证件背面",
    "registration_portrait": "二寸免冠照片",
    "portrait": "二寸免冠照片",
    "graduation_certificate": "学历证明",
    "education_proof": "学历证明",
    "education_report": "学信网电子注册表",
    "xuexin_registration": "学信网电子注册表",
    "commitment_letter": "承诺书",
    "student_card": "学生证",
    "social_security_proof": "社保证明",
    "vocational_qualification_certificate": "职业资格证书",
    "professional_title_certificate": "职称证书",
}
EMPLOYMENT_REPORTING = {
    "confirmed": True,
    "date_rule": "work_only",
    "years_rule": "required",
    "allow_zero_years": True,
}
RULE_ID = "renshe-degree-and-four-level-employment-20261018"


def archive_download_name(plan_id: int | str, volume_no: int) -> str:
    return f"人社报名_批次{plan_id}_第{volume_no}卷.zip"


def employment_rules() -> dict:
    return dict(EMPLOYMENT_REPORTING)


def work_date(info: dict) -> str:
    value = info.get("work_start_date")
    if not value:
        raise BusinessException(
            "请补齐真实参加工作日期；入学日期、毕业日期不能替代，未参加工作请联系管理员"
        )
    try:
        parsed = date.fromisoformat(str(value))
    except ValueError as exc:
        raise BusinessException("参加工作日期格式无效") from exc
    return f"{parsed.year}-{parsed.month}-{parsed.day}"


def material_archive_path(id_card: str, kind: str, extension: str) -> str:
    if not re.fullmatch(r"\d{17}[\dXx]", id_card):
        raise BusinessException("导出快照中的身份证号无效")
    if kind not in MATERIAL_STEMS or extension.lower() not in {".jpg", ".jpeg", ".pdf"}:
        raise BusinessException("导出材料类型或真实扩展名不受支持")
    return f"{id_card}_{MATERIAL_STEMS[kind]}{extension.lower()}"


def definition_fingerprint(registration: Path, work: Path) -> str:
    from app.services.renshe_reporting import HEADERS, DEGREE_CONDITION
    from app.services.renshe_export import build_registration_row, RensheExportService
    from app.services.renshe_catalog import combinations
    from app.services.renshe_employment import validate_qualification, cumulative_years

    content = {
        "rule": RULE_ID,
        "headers": HEADERS,
        "condition": DEGREE_CONDITION,
        "mapping": inspect.getsource(build_registration_row),
        "dictionary": inspect.getsource(prepare_reporting_workbook),
        "employment_validation": inspect.getsource(validate_qualification) + inspect.getsource(cumulative_years),
        "work_mapping": inspect.getsource(RensheExportService._write_work_history_workbook),
        "programs": [item for item in combinations() if item["adapted"]],
        "materials": MATERIAL_STEMS,
        "employment": employment_rules(),
        "templates": [
            hashlib.sha256(path.read_bytes()).hexdigest()
            for path in (registration, work)
        ],
    }
    return hashlib.sha256(
        json.dumps(content, sort_keys=True, ensure_ascii=False).encode()
    ).hexdigest()


def prepare_reporting_workbook(workbook) -> None:
    """Extend the dictionary on a private copy, retaining 4 sheets/25 columns.

    The named range avoids Excel's 255-character literal-list limit. U retains
    the client's OFFSET/MATCH link to each contiguous occupation's conditions.
    """
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.defined_name import DefinedName
    from app.services.renshe_catalog import combinations

    selected = [item for item in combinations() if item["adapted"]]
    conditions = workbook["申报条件"]
    existing = {(row[0], row[1]) for row in conditions.iter_rows(min_row=2, max_col=2, values_only=True)}
    codes = list(dict.fromkeys(item["occupation_code"] for item in selected))
    for item in selected:
        for condition in item["source_conditions"]:
            pair = (item["occupation_code"], condition)
            if pair not in existing:
                conditions.append(pair)
                existing.add(pair)
    conditions.cell(1, 4, "当前已适配报考职业代码")
    for i, code in enumerate(codes, 2):
        conditions.cell(i, 4, code)
    conditions.column_dimensions["D"].hidden = True
    workbook.defined_names.add(DefinedName("renshe_current_occupations", attr_text=f"'申报条件'!$D$2:$D${len(codes) + 1}"))
    sheet = workbook["Sheet1"]
    rules = sheet.data_validations.dataValidation
    for rule in list(rules):
        if rule.type == "list" and "T2" in rule.sqref:
            rules.remove(rule)
    rule = DataValidation(type="list", formula1="renshe_current_occupations", allow_blank=False)
    rule.errorTitle, rule.error = "报考职业不匹配", "请选择当前已适配的报考职业代码"
    rule.showErrorMessage = True
    sheet.add_data_validation(rule)
    rule.add("T2:T1048576")


def export_instructions(config: dict | None = None) -> dict:
    from app.services.renshe_reporting import HEADERS

    examples = [
        ("id_card_front", ".jpg"),
        ("id_card_back", ".jpeg"),
        ("registration_portrait", ".jpg"),
        ("graduation_certificate", ".pdf"),
        ("education_report", ".pdf"),
        ("commitment_letter", ".pdf"),
    ]
    if config is not None:
        examples = [
            (item["kind"], next((ext for ext in (".pdf", ".jpeg", ".jpg") if ext in item["extensions"]), ".jpg"))
            for item in config.get("materials", []) if item["kind"] in MATERIAL_STEMS
        ]
    masked = "110101********123X"
    paths = [REGISTRATION_TEMPLATE_NAME, WORK_HISTORY_TEMPLATE_NAME]
    paths += [
        material_archive_path("11010119900101123X", kind, ext).replace(
            "11010119900101123X", masked
        )
        for kind, ext in examples
    ]
    sources = [
        "本人实名姓名",
        "本人实名性别",
        "本人实名出生日期",
        "居民身份证",
        "本人实名身份证号",
        "社会人员",
        "本次报名学历（按模板文化程度转换）",
        "本次填报民族",
        "本次填报政治面貌",
        "本次联系电话",
        "本次通讯地址",
        "本次工作单位（选填）",
        "本人实际从事职业",
        "批次所选职业／工种",
        "批次所选技能等级的导出值",
        "学历型不填写职称专业",
        "学历型不填写职称等级",
        "真实参加工作日期，格式如2022-9-21；不使用入学或毕业日期",
        "实际职业年限，必填，允许0",
        "批次匹配的职业代码",
        "批次组合匹配的模板申报条件原文",
        "批次所选申报类型",
        "批次考试类型（新考／补考）",
        "实际社保参保单位（选填）",
        "本次电子邮箱",
    ]
    if config and config.get("qualification_rule") == "documented_employment_v1":
        sources[6] = "本次实际填报文化程度（不以毕业学历作为本路径门槛）"
        sources[11] = "本次实际工作单位（必填，与工龄证明核对）"
        sources[15] = "本路径不填写职称专业"
        sources[16] = "本路径不填写职称等级"
        sources[23] = "实际社保参保单位（必填，与社保证明核对）"
    tree = [
        archive_download_name("编号", 1),
        f"├─ {paths[0]}",
        f"├─ {paths[1]}",
    ]
    for i, path in enumerate(paths[2:]):
        tree.append(
            f"{'└─' if i == len(paths[2:]) - 1 else '├─'} {path.rsplit('/', 1)[-1]}"
        )
    return {
        "registration_template": REGISTRATION_TEMPLATE_NAME,
        "work_template": WORK_HISTORY_TEMPLATE_NAME,
        "work_description": "按本次报名版本的实际工作经历逐条生成；结束日期为空表示至今。证书及其网查截图合并为同一份JPEG材料，工龄证明独立导出。" if config and config.get("qualification_rule") == "documented_employment_v1" else "原样附带甲方工作经历模板，目前不会按考生自动填写经历；后续按报送要求补充。",
        "directory_note": "目录示例，实际内容以报名材料为准；文件保留真实扩展名。",
        "portrait_note": "本期报名照片为一寸白底近期证件照、≤500 KB；文件名沿用甲方既定的二寸免冠照片，不改变真实格式。历史批次保留原材料要求。",
        "example_paths": paths,
        "file_tree": "\n".join(tree),
        "columns": [
            {"column": i, "header": h, "source": s}
            for i, (h, s) in enumerate(zip(HEADERS, sources), 1)
        ],
    }
