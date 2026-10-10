"""NISP async export service with job queue, OSS storage, and signed URLs."""

import re
import tempfile
import zipfile
import hashlib
import uuid
from pathlib import Path
from datetime import datetime, timedelta, timezone

from sqlalchemy import func, select, update as sa_update

from app.adapter.database import get_db_ctx
from app.domain.nisp.src import (
    NispExportItem,
    NispExportJob,
    NispMaterialFile,
    NispRegistration,
    NispRegistrationVersion,
)
from app.integrations.nisp_storage import NispObjectStorage
from app.port.exceptions import BusinessException, NotFoundException
from app.schemas.nisp import NispExportJobResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NISP_LEVEL1_HEADERS = [
    "序号", "姓名", "拼音", "专业", "学校名称/单位名称",
    "身份证号(18位)", "手机号码", "邮箱", "报考省份", "培训种类",
    "身份证双面", "寸照",
]

NISP_LEVEL2_HEADERS = [
    "序号", "培训机构", "姓名", "拼音", "学校名称",
    "身份证号(18位)", "考试类型", "性别", "年龄", "最高学历",
    "专业", "移动电话", "邮箱", "省份", "地址", "邮编",
    "身份证双面照片", "寸照", "学籍报告", "NISP二级考试报名申请表",
]

HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
HEADER_FONT = Font(name="微软雅黑", size=10, bold=True)
BODY_FONT = Font(name="微软雅黑", size=10)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

EXPORT_TTL_HOURS = 72
WORKER_POLL_SECONDS = 5
REQUIRED_MATERIALS = {
    "1": ("id_card_both_sides", "portrait_photo"),
    "2": (
        "id_card_both_sides",
        "portrait_photo",
        "xuexin_report",
        "application_form",
    ),
}
MATERIAL_EXPORT_NAMES = {
    "id_card_both_sides": "身份证双面",
    "portrait_photo": "寸照",
    "xuexin_report": "学籍报告",
    "application_form": "NISP二级考试报名申请表",
}
MANIFEST_HEADERS = [
    "报名编号", "姓名", "身份证号", "个人压缩包", "身份证双面", "寸照",
    "学籍报告", "NISP二级考试报名申请表",
]


def _now() -> datetime:
    return datetime.now(timezone.utc)


class NispExportIncompleteMaterials(Exception):
    def __init__(self, missing: list[dict]):
        self.missing = missing
        super().__init__("部分审核通过报名的材料不完整")


class NispExportService:

    def __init__(self, storage: NispObjectStorage | None = None) -> None:
        self.storage = storage or NispObjectStorage()

    async def create_job(
        self,
        *,
        admin_id: int,
        batch_id: int,
        level: str,
        include_statuses: list[str] | None = None,
    ) -> NispExportJobResponse:
        include_statuses = include_statuses or ["approved"]
        async with get_db_ctx() as db:
            registrations = (
                await db.execute(
                    select(NispRegistration)
                    .where(
                    NispRegistration.batch_id == batch_id,
                    NispRegistration.level == level,
                    NispRegistration.status.in_(include_statuses),
                )
                    .order_by(NispRegistration.id)
                    .with_for_update()
                )
            ).scalars().all()
            if not registrations:
                raise BusinessException("没有符合条件的报名记录")

            job = NispExportJob(
                batch_id=batch_id,
                level=level,
                requested_by_admin_id=admin_id,
                include_statuses=include_statuses,
                status="queued",
                artifact_type="full_package",
                registration_count=len(registrations),
            )
            db.add(job)
            await db.flush()
            for registration in registrations:
                version = await db.scalar(
                    select(NispRegistrationVersion)
                    .where(
                        NispRegistrationVersion.registration_id == registration.id,
                        NispRegistrationVersion.is_current.is_(True),
                    )
                    .with_for_update()
                )
                if version is None:
                    raise BusinessException("NISP 报名缺少当前信息版本")
                db.add(NispExportItem(
                    job_id=job.id,
                    registration_id=registration.id,
                    registration_version_id=version.id,
                    candidate_snapshot=version.candidate_snapshot,
                    material_versions=version.material_versions,
                    is_valid=True,
                ))
            await db.commit()
            await db.refresh(job)
            return NispExportJobResponse.model_validate(job)

    async def list_jobs(
        self, *, page: int = 1, page_size: int = 20
    ) -> tuple[list[NispExportJobResponse], int]:
        async with get_db_ctx() as db:
            total = await db.scalar(
                select(func.count()).select_from(NispExportJob)
            ) or 0
            rows = (
                await db.execute(
                    select(NispExportJob)
                    .order_by(NispExportJob.id.desc())
                    .offset((page - 1) * page_size)
                    .limit(page_size)
                )
            ).scalars().all()
            return [NispExportJobResponse.model_validate(r) for r in rows], total

    async def get_job(self, job_id: int) -> NispExportJobResponse:
        async with get_db_ctx() as db:
            job = await db.get(NispExportJob, job_id)
            if job is None:
                raise NotFoundException("NISP 导出任务")
            return NispExportJobResponse.model_validate(job)

    async def signed_url(self, job_id: int) -> str:
        async with get_db_ctx() as db:
            job = await db.get(NispExportJob, job_id)
            if job is None:
                raise NotFoundException("NISP 导出任务")
            if job.status != "succeeded" or not job.storage_key:
                raise BusinessException("导出任务尚未完成")
            if job.expires_at and job.expires_at < _now():
                raise BusinessException("导出文件已过期，请重新导出")
            return await self.storage.signed_get_url(job.storage_key)

    async def process_next_job(self) -> bool:
        """Pick up and process one queued job. Returns True if a job was processed."""
        async with get_db_ctx() as db:
            await db.execute(
                sa_update(NispExportJob)
                .where(
                    NispExportJob.status == "running",
                    NispExportJob.lease_expires_at < _now(),
                )
                .values(
                    status="queued",
                    last_error="worker lease expired; job recovered",
                )
            )
            await db.commit()
            job = (
                await db.execute(
                    select(NispExportJob)
                    .where(NispExportJob.status == "queued")
                    .order_by(NispExportJob.id.asc())
                    .limit(1)
                    .with_for_update(skip_locked=True)
                )
            ).scalar_one_or_none()
            if job is None:
                return False

            job.status = "running"
            job.started_at = _now()
            job.heartbeat_at = _now()
            job.lease_expires_at = _now() + timedelta(minutes=10)
            await db.commit()

            try:
                await self._run_job(job.id)
            except NispExportIncompleteMaterials as exc:
                async with get_db_ctx() as db2:
                    failed = await db2.get(NispExportJob, job.id)
                    if failed is not None:
                        failed.status = "failed"
                        failed.last_error = str(exc)
                        failed.finished_at = _now()
                        failed.result_summary = {
                            "missing_materials": exc.missing,
                            "missing_count": len(exc.missing),
                        }
                        await db2.commit()
            except Exception as exc:
                async with get_db_ctx() as db2:
                    failed = await db2.get(NispExportJob, job.id)
                    if failed is not None:
                        failed.status = "failed"
                        failed.last_error = f"{type(exc).__name__}: {exc}"[:2000]
                        failed.finished_at = _now()
                        await db2.commit()
            return True

    async def _run_job(self, job_id: int) -> None:
        async with get_db_ctx() as db:
            job = await db.get(NispExportJob, job_id)
            if job is None:
                return
            registrations = (
                await db.execute(
                    select(NispRegistration)
                    .join(
                        NispExportItem,
                        NispExportItem.registration_id == NispRegistration.id,
                    )
                    .join(
                        NispRegistrationVersion,
                        NispRegistrationVersion.id
                        == NispExportItem.registration_version_id,
                    )
                    .where(
                        NispExportItem.job_id == job.id,
                        NispExportItem.is_valid.is_(True),
                        NispRegistrationVersion.is_current.is_(True),
                    )
                    .order_by(NispRegistration.id.asc())
                )
            ).scalars().all()
            if not registrations:
                raise BusinessException("NISP 导出清单中没有有效报名版本")

            material_rows = (
                await db.execute(
                    select(NispMaterialFile)
                    .where(
                        NispMaterialFile.registration_id.in_(
                            [item.id for item in registrations]
                        ),
                        NispMaterialFile.is_current.is_(True),
                    )
                    .order_by(NispMaterialFile.registration_id, NispMaterialFile.id)
                )
            ).scalars().all()
            materials_by_registration: dict[int, dict[str, NispMaterialFile]] = {}
            for material in material_rows:
                materials_by_registration.setdefault(material.registration_id, {})[
                    material.material_type
                ] = material

            missing = []
            required = REQUIRED_MATERIALS[job.level]
            for reg in registrations:
                current = materials_by_registration.get(reg.id, {})
                missing_types = [item for item in required if item not in current]
                if missing_types:
                    missing.append(
                        {
                            "registration_id": reg.id,
                            "registration_no": reg.registration_no,
                            "name": reg.candidate_snapshot.get("name", ""),
                            "missing_materials": missing_types,
                        }
                    )
            if missing:
                raise NispExportIncompleteMaterials(missing)

            manifests = self._manifests(
                list(registrations), materials_by_registration
            )
            excel_bytes = self._build_excel(
                list(registrations),
                level=job.level,
                manifests=manifests,
            )
            storage_key = f"nisp/exports/{job.id}/{uuid.uuid4().hex}.zip"
            with tempfile.TemporaryDirectory(
                prefix=f"nisp-export-upload-{job.id}-"
            ) as upload_raw:
                package_path = await self._build_full_package(
                    job_id=job.id,
                    level=job.level,
                    excel_bytes=excel_bytes,
                    registrations=list(registrations),
                    materials_by_registration=materials_by_registration,
                    manifests=manifests,
                    destination=Path(upload_raw) / "full-package.zip",
                )
                await self.storage.upload_export_file(
                    storage_key, package_path, "application/zip"
                )
                artifact_bytes, artifact_sha256 = self._file_size_and_sha256(
                    package_path
                )

            job.status = "succeeded"
            job.finished_at = _now()
            job.storage_key = storage_key
            job.artifact_type = "full_package"
            job.artifact_sha256 = artifact_sha256
            job.artifact_bytes = artifact_bytes
            job.expires_at = _now() + timedelta(hours=EXPORT_TTL_HOURS)
            job.result_summary = {
                "registration_count": len(registrations),
                "package_count": len(registrations),
                "artifact_type": "full_package",
                "packages": [
                    {
                        "registration_no": reg.registration_no,
                        "name": reg.candidate_snapshot.get("name", ""),
                        "filename": manifests[reg.id]["package"],
                    }
                    for reg in registrations
                ],
            }
            await db.commit()

    async def expire_artifacts(self) -> int:
        """Mark expired succeeded jobs as failed and delete their artifacts."""
        now = _now()
        async with get_db_ctx() as db:
            jobs = (
                await db.execute(
                    select(NispExportJob)
                    .where(
                        NispExportJob.status == "succeeded",
                        NispExportJob.expires_at < now,
                    )
                )
            ).scalars().all()
            for job in jobs:
                if job.storage_key:
                    try:
                        await self.storage.delete_object(job.storage_key)
                    except Exception:
                        pass
                job.status = "failed"
                job.last_error = "expired"
                job.storage_key = None
            await db.commit()
            return len(jobs)

    def _build_excel(
        self,
        registrations: list[NispRegistration],
        *,
        level: str,
        manifests: dict[int, dict[str, str]] | None = None,
    ) -> bytes:
        import io

        wb = Workbook()
        if level == "1":
            ws = wb.active
            ws.title = "NISP一级报名表"
            self._write_level1(ws, registrations)
        else:
            ws = wb.active
            ws.title = "NISP二级报名表"
            self._write_level2(ws, registrations)

        if manifests is not None:
            manifest_ws = wb.create_sheet("资料清单")
            self._write_manifest(manifest_ws, registrations, manifests)

        headers = NISP_LEVEL1_HEADERS if level == "1" else NISP_LEVEL2_HEADERS
        for col_idx in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 16

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.read()

    def _manifests(
        self,
        registrations: list[NispRegistration],
        materials_by_registration: dict[int, dict[str, NispMaterialFile]],
    ) -> dict[int, dict[str, str]]:
        name_counts: dict[str, int] = {}
        for reg in registrations:
            name = self._safe_filename(str(reg.candidate_snapshot.get("name", "")))
            if not name:
                name = f"报名{reg.registration_no}"
            name_counts[name] = name_counts.get(name, 0) + 1

        seen_names: dict[str, int] = {}
        result: dict[int, dict[str, str]] = {}
        for reg in registrations:
            name = self._safe_filename(str(reg.candidate_snapshot.get("name", "")))
            if not name:
                name = f"报名{reg.registration_no}"
            seen_names[name] = seen_names.get(name, 0) + 1
            package_name = (
                f"{name}-{reg.registration_no}.zip"
                if seen_names[name] > 1
                else f"{name}.zip"
            )
            materials = materials_by_registration[reg.id]
            id_card_name = f"{name}.pdf"
            portrait_name = f"{name}-{reg.candidate_idcard}.jpg"
            xuexin_name = f"{name}-学籍报告.pdf"
            application = materials.get("application_form")
            application_name = ""
            if application is not None:
                application_name = self._safe_filename(
                    application.original_filename
                    or MATERIAL_EXPORT_NAMES["application_form"]
                )
                if not application_name.lower().endswith(".pdf"):
                    application_name += ".pdf"
            inner_names = {id_card_name, portrait_name, xuexin_name}
            if application_name and application_name in inner_names:
                stem = application_name[:-4]
                candidate = f"{stem}-NISP二级考试报名申请表.pdf"
                serial = 2
                while candidate in inner_names:
                    candidate = (
                        f"{stem}-NISP二级考试报名申请表-{serial}.pdf"
                    )
                    serial += 1
                application_name = candidate
            result[reg.id] = {
                "name": name,
                "package": package_name,
                "id_card_both_sides": id_card_name,
                "portrait_photo": portrait_name,
                "xuexin_report": xuexin_name,
                "application_form": application_name,
            }
        return result

    async def _build_full_package(
        self,
        *,
        job_id: int,
        level: str,
        excel_bytes: bytes,
        registrations: list[NispRegistration],
        materials_by_registration: dict[int, dict[str, NispMaterialFile]],
        manifests: dict[int, dict[str, str]],
        destination: Path,
    ) -> Path:
        with tempfile.TemporaryDirectory(prefix=f"nisp-export-{job_id}-") as raw:
            root = Path(raw)
            staging = root / "materials"
            staging.mkdir()
            destination.parent.mkdir(parents=True, exist_ok=True)
            excel_name = (
                "NISP一级报名汇总表.xlsx" if level == "1" else "NISP二级报名汇总表.xlsx"
            )
            with zipfile.ZipFile(
                destination, mode="w", compression=zipfile.ZIP_DEFLATED, allowZip64=True
            ) as outer:
                outer.writestr(excel_name, excel_bytes)
                for reg in registrations:
                    manifest = manifests[reg.id]
                    person_path = root / "packages" / manifest["package"]
                    person_path.parent.mkdir(parents=True, exist_ok=True)
                    with zipfile.ZipFile(
                        person_path,
                        mode="w",
                        compression=zipfile.ZIP_DEFLATED,
                        allowZip64=True,
                    ) as person:
                        for material_type in REQUIRED_MATERIALS[level]:
                            material = materials_by_registration[reg.id][material_type]
                            material_path = staging / f"{material.id}-{uuid.uuid4().hex}"
                            await self.storage.download_file(
                                material.storage_key, material_path
                            )
                            person.write(
                                material_path,
                                arcname=manifest[material_type],
                            )
                            material_path.unlink(missing_ok=True)
                    outer.write(person_path, arcname=manifest["package"])
                    person_path.unlink(missing_ok=True)
            return destination

    def _write_manifest(
        self,
        ws,
        registrations: list[NispRegistration],
        manifests: dict[int, dict[str, str]],
    ) -> None:
        ws.append(MANIFEST_HEADERS)
        self._style_header(ws, len(MANIFEST_HEADERS))
        for reg in registrations:
            manifest = manifests[reg.id]
            ws.append(
                [
                    reg.registration_no,
                    manifest["name"],
                    reg.candidate_idcard,
                    manifest["package"],
                    manifest["id_card_both_sides"],
                    manifest["portrait_photo"],
                    manifest.get("xuexin_report", "-"),
                    manifest.get("application_form", "-"),
                ]
            )
            self._style_row(ws, ws.max_row, len(MANIFEST_HEADERS))
        for col_idx in range(1, len(MANIFEST_HEADERS) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 22

    @staticmethod
    def _safe_filename(value: str) -> str:
        normalized = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', "", value).strip().strip(".")
        return normalized[:180]

    @staticmethod
    def _file_size_and_sha256(path: Path) -> tuple[int, str]:
        digest = hashlib.sha256()
        size = 0
        with path.open("rb") as source:
            while chunk := source.read(1024 * 1024):
                size += len(chunk)
                digest.update(chunk)
        return size, digest.hexdigest()

    def _write_level1(self, ws, registrations: list[NispRegistration]) -> None:
        ws.append(NISP_LEVEL1_HEADERS)
        self._style_header(ws, len(NISP_LEVEL1_HEADERS))
        for idx, reg in enumerate(registrations, 1):
            s = reg.candidate_snapshot
            row = [
                idx, s.get("name", ""), s.get("pinyin", ""), s.get("major", ""),
                s.get("school", ""), s.get("id_card", reg.candidate_idcard),
                s.get("phone", ""), s.get("email", ""), s.get("province", ""),
                s.get("training_type", "NISP一级"),
                "已上传" if self._has_material(reg, "id_card_both_sides") else "未上传",
                "已上传" if self._has_material(reg, "portrait_photo") else "未上传",
            ]
            ws.append(row)
            self._style_row(ws, ws.max_row, len(NISP_LEVEL1_HEADERS))

    def _write_level2(self, ws, registrations: list[NispRegistration]) -> None:
        ws.append(NISP_LEVEL2_HEADERS)
        self._style_header(ws, len(NISP_LEVEL2_HEADERS))
        for idx, reg in enumerate(registrations, 1):
            s = reg.candidate_snapshot
            row = [
                idx, s.get("institution", ""), s.get("name", ""), s.get("pinyin", ""),
                s.get("school", ""), s.get("id_card", reg.candidate_idcard),
                s.get("training_type", "NISP二级"), s.get("gender", ""),
                s.get("age", ""), s.get("education", ""), s.get("major", ""),
                s.get("phone", ""), s.get("email", ""), s.get("province", ""),
                s.get("address", ""), s.get("zip_code", ""),
                "已上传" if self._has_material(reg, "id_card_both_sides") else "未上传",
                "已上传" if self._has_material(reg, "portrait_photo") else "未上传",
                "已上传" if self._has_material(reg, "xuexin_report") else "未上传",
                "已上传" if self._has_material(reg, "application_form") else "未上传",
            ]
            ws.append(row)
            self._style_row(ws, ws.max_row, len(NISP_LEVEL2_HEADERS))

    @staticmethod
    def _has_material(reg: NispRegistration, material_type: str) -> bool:
        keys = reg.material_keys or {}
        return bool(keys.get(material_type))

    @staticmethod
    def _style_header(ws, col_count: int) -> None:
        for col in range(1, col_count + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = THIN_BORDER
            cell.alignment = CENTER_ALIGN
        ws.row_dimensions[1].height = 28

    @staticmethod
    def _style_row(ws, row: int, col_count: int) -> None:
        for col in range(1, col_count + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = CENTER_ALIGN


async def nisp_export_worker_loop(service: NispExportService | None = None) -> None:
    import asyncio
    import logging

    logger = logging.getLogger(__name__)
    active_service = service or NispExportService()
    while True:
        try:
            processed = await active_service.process_next_job()
            if not processed:
                await active_service.expire_artifacts()
        except Exception:
            logger.warning("NISP export worker iteration failed", exc_info=True)
        await asyncio.sleep(WORKER_POLL_SECONDS)
