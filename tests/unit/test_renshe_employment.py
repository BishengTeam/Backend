"""Work and certificate paths: boundaries, real SQL/storage and report packages."""
import io
import zipfile
from datetime import UTC, date, datetime, timedelta

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select

from app.domain.order.src.index import Order
from app.domain.plan.src.index import Plan
from app.domain.renshe.src.index import RensheExportVolume
from app.port.exceptions import BusinessException
from app.schemas.renshe import RensheDraftUpsert, RensheReviewCreate
from app.schemas.renshe_project import RensheProjectWrite
from app.services.order_fulfillment import OrderFulfillmentService
from app.services.plan import PlanService
from app.services.renshe_application import RensheApplicationService
from app.services.renshe_catalog import ADAPTED_DEGREE_OCCUPATIONS
from app.services.renshe_employment import (
    anniversary,
    cumulative_years,
    review_checks,
    validate_qualification,
)
from app.services.renshe_export import RensheExportService
from app.services.renshe_graduated import china_today
from app.services.renshe_project import RensheProjectService
from app.services.renshe_review import RensheReviewService
from tests.unit.test_renshe_graduated import graduated_batch
from tests.unit.test_renshe_projects import create_batch
from tests.unit.test_renshe_projects import isolated as _isolated

isolated = _isolated


def payload(batch, *, compound=False):
    today = china_today()
    start = anniversary(today, -6).isoformat()
    return RensheDraftUpsert(
        plan_id=batch.id, contact_phone="13800138000", email="work@example.test", mailing_address="真实通讯地址",
        reporting_info={"occupation_status": "职工", "ethnicity": "汉族", "political_status": "群众", "has_work_experience": True,
                        "work_start_date": start, "occupation_years": 6, "work_unit": "样例单位", "social_insurance_unit": "样例单位"},
        qualification_info={"cultural_education": "primary", "work_experiences": [{"work_unit": "样例单位", "position": "实际岗位", "occupation_name": batch.occupation_name,
                             "start_date": start, "end_date": None}],
                            **({"skill_certificate": {"occupation_name": batch.occupation_name, "grade": 5, "certificate_number": "TEST-L5", "issued_on": anniversary(today, -4).isoformat()}} if compound else {})},
    )


async def employment_batch(env, monkeypatch, occupation="information_security", application_type="工龄型"):
    original, _ = await graduated_batch(env, monkeypatch)
    project = await RensheProjectService().save(RensheProjectWrite(
        chinese_name=f"{occupation}-{application_type}", selection_id=f"{occupation}:4:{application_type}",
        commitment_template=original.config.commitment_template, status="available",
    ), admin_id=1)
    batch = await create_batch(project)
    await PlanService().publish_plan(batch.id, project.code, admin_id=1)
    service = RensheApplicationService(env.storage)
    for item in batch.renshe_config["materials"]:
        kind = item["kind"]
        if kind.startswith("id_card_"):
            continue
        pdf = kind in {"work_proof", "social_security_proof", "commitment_letter"}
        await service.upload_application_material(1, batch.id, kind, filename="合并原件.jpg" if not pdf else "证明.pdf", content_type="application/pdf" if pdf else "image/jpeg", content=b"%PDF-1.4\nfixture" if pdf else b"\xff\xd8\xfffixture")
    return project, batch, service


async def pay(env, submitted):
    async with env.db() as db:
        order = await db.get(Order, submitted.order_id)
        order.status, order.paid_at, order.transaction_id = "paid", datetime.now(UTC), f"work-{order.id}"
        await OrderFulfillmentService().on_paid(db, order)
        await db.commit()


@pytest.mark.asyncio
@pytest.mark.parametrize("occupation", sorted(ADAPTED_DEGREE_OCCUPATIONS))
@pytest.mark.parametrize("application_type", ["工龄型", "复合型"])
async def test_each_trade_submits_reviews_and_exports_work_path(isolated, monkeypatch, occupation, application_type):
    _, batch, service = await employment_batch(isolated, monkeypatch, occupation, application_type)
    data = payload(batch, compound=application_type == "复合型")
    draft = await service.save_draft(1, data)
    submitted = await service.submit(1, draft.id)
    await pay(isolated, submitted)
    reviewer = RensheReviewService()
    with pytest.raises(BusinessException, match="逐项"):
        await reviewer.review(reviewer_id=1, application_id=draft.id, stage="initial", data=RensheReviewCreate(decision="approved", expected_version_id=submitted.version.id))
    await reviewer.review(reviewer_id=1, application_id=draft.id, stage="initial", data=RensheReviewCreate(decision="approved", expected_version_id=submitted.version.id, checklist={key: True for key in review_checks(batch.renshe_config)}))
    exporter = RensheExportService(isolated.storage)
    job = await exporter.create_job(plan_id=batch.id, admin_id=1)
    await exporter.process_next_job()
    finished = await exporter.get_job(job.id)
    assert finished.status == "succeeded", finished.last_error
    async with isolated.db() as db:
        volume = await db.scalar(select(RensheExportVolume).where(RensheExportVolume.job_id == job.id))
        archive_path = isolated.storage._local_path(volume.storage_key)
    with zipfile.ZipFile(archive_path) as archive:
        files = archive.namelist()
        assert any(name.endswith("_工龄证明.pdf") for name in files)
        assert any(name.endswith("_社保证明.pdf") for name in files)
        assert not any("学历证明" in name or "网查截图" in name for name in files)
        assert sum(name.endswith("_职业资格证书.jpg") for name in files) == (1 if application_type == "复合型" else 0)
        work = load_workbook(io.BytesIO(archive.read("工作经历.xlsx")))
        assert work["Sheet1"].max_row == 2
        assert work["Sheet1"]["C2"].value == data.qualification_info.work_experiences[0].start_date.isoformat()
        assert work["Sheet1"]["D2"].value is None
        assert work["Sheet1"]["E2"].value == "样例单位" and work["Sheet1"]["F2"].value == "实际岗位"
        work.close()
        registration = load_workbook(io.BytesIO(archive.read("报名信息.xlsx")))
        assert registration["Sheet1"]["G2"].value == "小学"
        assert registration["Sheet1"]["N2"].value == batch.occupation_name
        assert registration["Sheet1"]["V2"].value == application_type
        assert registration["Sheet1"]["S2"].value == 6
        registration.close()
    await reviewer.review(reviewer_id=1, application_id=draft.id, stage="external", data=RensheReviewCreate(decision="approved", expected_version_id=submitted.version.id))


def test_calendar_anniversaries_gaps_and_overlap_are_not_counted_twice():
    row = {"start_date": "2020-01-01", "end_date": "2024-12-31"}
    assert cumulative_years([row, row], as_of=date(2025, 1, 1)) == 5
    assert cumulative_years([{"start_date": "2020-02-29", "end_date": "2025-02-27"}], as_of=date(2025, 2, 28)) == 5
    assert cumulative_years([row], as_of=date(2025, 1, 1), after=date(2023, 1, 1)) == 2
    assert cumulative_years([{"start_date": "2020-01-01", "end_date": "2021-12-31"}, {"start_date": "2023-01-01", "end_date": "2024-12-31"}], as_of=date(2025, 1, 1)) == 4


@pytest.mark.asyncio
async def test_paid_work_correction_reuses_order_and_copies_new_snapshot(isolated, monkeypatch):
    _, batch, service = await employment_batch(isolated, monkeypatch)
    draft = await service.save_draft(1, payload(batch))
    submitted = await service.submit(1, draft.id)
    await pay(isolated, submitted)
    await RensheReviewService().review(reviewer_id=1, application_id=draft.id, stage="initial", data=RensheReviewCreate(
        decision="rejected", expected_version_id=submitted.version.id, reason="请核对单位和岗位", required_changes=["qualification_info.work_experiences", "work_proof"], correction_due_at=datetime.now(UTC) + timedelta(days=3)))
    async with isolated.db() as db:
        stored = await db.get(Plan, batch.id)
        stored.status = "registration_closed"
        stored.apply_end = datetime.now(UTC) - timedelta(days=1)
        await db.commit()
    await service.save_draft(1, payload(batch))
    resubmitted = await service.submit(1, draft.id)
    assert resubmitted.order_id == submitted.order_id and not resubmitted.payment_required
    assert resubmitted.version.version_no == 2


@pytest.mark.parametrize("compound", [False, True])
def test_short_or_certificate_prior_work_and_missing_social_unit_are_rejected(compound):
    from app.services.renshe_catalog import find_selection
    from app.services.renshe_project import selection_configuration
    config = selection_configuration(find_selection(f"information_security:4:{'复合型' if compound else '工龄型'}"))
    batch = type("Batch", (), {"id": 1, "occupation_name": "信息安全管理员"})()
    form = payload(batch, compound=compound).model_dump(mode="json")
    if compound:
        form["qualification_info"]["skill_certificate"]["issued_on"] = anniversary(china_today(), -2).isoformat()
    else:
        form["qualification_info"]["work_experiences"][0]["start_date"] = anniversary(china_today(), -4).isoformat()
        form["reporting_info"]["work_start_date"] = form["qualification_info"]["work_experiences"][0]["start_date"]
    with pytest.raises(BusinessException, match="累计"):
        validate_qualification(form, config)
    form = payload(batch, compound=compound).model_dump(mode="json")
    form["reporting_info"]["social_insurance_unit"] = ""
    with pytest.raises(BusinessException, match="社保"):
        validate_qualification(form, config)


def test_client_cannot_supply_assessment_snapshot_or_unknown_nested_fields():
    batch = type("Batch", (), {"id": 1, "occupation_name": "信息安全管理员"})()
    form = payload(batch).model_dump(mode="json")
    with pytest.raises(ValidationError):
        RensheDraftUpsert.model_validate({**form, "eligibility_snapshot": {"as_of": "2099-01-01"}})
