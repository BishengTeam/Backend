"""Fourth-level source alignment, independent thresholds and material boundaries."""

from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.port.exceptions import BusinessException
from app.services.renshe_catalog import ADAPTED_DEGREE_OCCUPATIONS, combinations, find_selection
from app.services.renshe_employment import anniversary
from app.services.renshe_graduated import china_today, validate_admission, review_checks
from app.services.renshe_project import selection_configuration
from tests.unit.test_renshe_employment import payload


def test_fourth_level_has_exactly_twenty_one_source_backed_combinations():
    selected = [item for item in combinations() if item["grade"] == 4 and item["adapted"]]
    assert len(selected) == 21
    assert {item["occupation_key"] for item in selected} == ADAPTED_DEGREE_OCCUPATIONS
    workbook = load_workbook(Path(__file__).parents[2] / "docs/renshe/报名信息.xlsx")
    try:
        source = {row[1] for row in workbook["申报条件"].iter_rows(values_only=True) if row[0] == "4-04-04-02-网络与信息安全管理员-信息安全管理员-中级工"}
        assert {condition for item in selected for condition in item["source_conditions"]} == source
    finally:
        workbook.close()


@pytest.mark.parametrize("occupation", sorted(ADAPTED_DEGREE_OCCUPATIONS))
@pytest.mark.parametrize("application_type", ["工龄型", "复合型", "学历型"])
def test_fourth_level_materials_and_review_match_actual_selected_path(occupation, application_type):
    config = selection_configuration(find_selection(f"{occupation}:4:{application_type}"))
    kinds = {item["kind"] for item in config["materials"] if item["required"]}
    assert config["export_skill_level"] == "四级" and config["application_types"] == [application_type]
    assert {"id_card_front", "id_card_back", "registration_portrait", "commitment_letter"} <= kinds
    assert "student_card" not in kinds
    if application_type == "学历型":
        assert {"graduation_certificate", "education_report"} <= kinds
        assert not kinds & {"work_proof", "social_security_proof", "vocational_qualification_certificate"}
        assert {"graduated", "major", "report"} <= set(review_checks(config))
    else:
        assert {"work_proof", "social_security_proof"} <= kinds
        assert ("vocational_qualification_certificate" in kinds) == (application_type == "复合型")
        assert not kinds & {"graduation_certificate", "education_report"}
        assert ("满3年" if application_type == "复合型" else "满5年") in review_checks(config)["employment"]


@pytest.mark.parametrize("compound", [False, True])
def test_fourth_level_does_not_accept_other_grade_or_pre_certificate_work(compound):
    config = selection_configuration(find_selection(f"information_security:4:{'复合型' if compound else '工龄型'}"))
    data = payload(SimpleNamespace(id=1, occupation_name="信息安全管理员"), compound=compound).model_dump(mode="json")
    assert validate_admission(data, config)["eligibility_snapshot"]["required_years"] == (3 if compound else 5)
    if compound:
        data["qualification_info"]["skill_certificate"]["issued_on"] = anniversary(china_today(), -2).isoformat()
    else:
        data["qualification_info"]["work_experiences"][0]["start_date"] = anniversary(china_today(), -2).isoformat()
        data["qualification_info"]["work_experiences"] *= 3
    with pytest.raises(BusinessException, match="累计.*工作须满"):
        validate_admission(data, config)
    config["rule_profile"]["grade"] = 3
    with pytest.raises(BusinessException, match="独立资格规则不一致"):
        validate_admission(data, config)
