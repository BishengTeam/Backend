from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.port.exceptions import BusinessException
from app.schemas.sangfor import SangforCreate, SangforExport
from app.services.sangfor import export_workbook
from app.services.sangfor_catalog import DIRECTIONS, HEADERS, require_direction


def submission(**changes):
    return dict(batch_id=1, phone="13800000000", organization="测试学院", first_name="Xiaofu", last_name="Shen", exam_voucher_code="VENDOR-001") | changes


def test_only_unambiguous_vendor_directions_are_open():
    assert set(DIRECTIONS) == {"SCSA-S", "SCSA-T", "SCCA-C", "SCSP-S", "SCCP-C", "SCSE-T", "SCSE-S", "SCCE-C"}
    for code in ("SCSP-T", "SCSP-SOP", "SCSP-SOF", "SCCP-ST", "SCCP-S", "SCCA-LB", "SCCP-LB", "micro"):
        with pytest.raises(BusinessException, match="尚未确认"):
            require_direction(code)


@pytest.mark.parametrize("key", ["organization", "first_name", "last_name", "exam_voucher_code"])
def test_required_vendor_fields_reject_whitespace(key):
    with pytest.raises(ValidationError):
        SangforCreate(**submission(**{key: "   "}))


@pytest.mark.parametrize("changes", [{"phone": "138"}, {"first_name": "小服"}, {"email": "bad"}, {"candidate_name": "他人"}, {"verification_code": "fake"}])
def test_new_submission_rejects_invalid_or_untrusted_fields(changes):
    with pytest.raises(ValidationError):
        SangforCreate(**submission(**changes))


def test_optional_contact_fields_and_no_student_documents():
    value = SangforCreate(**submission())
    assert value.email == value.mailing_address == ""
    assert not {"student_proof", "verification_code", "coupon_code"} & set(value.model_dump())


def test_export_is_actual_thirteen_columns_with_typed_date_text_ids_and_formula_safe_voucher():
    snapshot = SangforCreate(**submission(exam_voucher_code="=2+2")).model_dump(exclude={"batch_id"}) | {
        "name": "实名考生", "id_card": "11010519491231002X", "exam_date": "2026-07-28T02:00:00+00:00", "exam_direction": DIRECTIONS["SCSA-T"]}
    book = load_workbook(BytesIO(export_workbook([snapshot])))
    sheet = book["Sheet1"]
    assert [sheet.cell(2, column).value for column in range(1, 14)] == HEADERS
    assert sheet["A3"].value == "实名考生"
    assert sheet["B3"].value == "11010519491231002X" and sheet["B3"].data_type == "s"
    assert sheet["D3"].value == datetime(2026, 7, 28, 10)
    assert sheet["D3"].number_format == "yyyy-mm-dd"
    assert sheet["H3"].value == "Xiaofu" and sheet["I3"].value == "Shen"
    assert sheet["J3"].value == DIRECTIONS["SCSA-T"]
    assert sheet["K3"].value == "=2+2" and sheet["K3"].data_type == "s"
    assert sheet["A4"].value is None


@pytest.mark.parametrize("ids", [[], [1, 1], [0], [-1]])
def test_export_selection_must_be_explicit_valid_unique(ids):
    with pytest.raises(ValidationError):
        SangforExport(registration_ids=ids)


def test_sangfor_http_only_opens_public_batches(monkeypatch):
    from types import SimpleNamespace
    from unittest.mock import AsyncMock
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.adapter.database import get_db
    from app.api.sangfor import router, service
    from app.middleware.error_handler import app_exception_handler
    from app.port.exceptions import AppException
    app = FastAPI()
    app.include_router(router, prefix="/api")
    app.add_exception_handler(AppException, app_exception_handler)
    async def db():
        yield SimpleNamespace()
    app.dependency_overrides[get_db] = db
    monkeypatch.setattr(service, "batches", AsyncMock(return_value=[]))
    client = TestClient(app)
    assert client.get("/api/sangfor/batches").status_code == 200
    assert client.get("/api/sangfor/registrations").status_code == 401
    assert client.get("/api/sangfor/registrations/1").status_code == 401
    assert client.post("/api/sangfor/registrations", json=submission()).status_code == 401
    assert client.post("/api/sangfor/registrations/1/resubmit", json={}).status_code == 401


def test_admin_refunds_uses_order_list_permission_and_validates_filters(monkeypatch):
    from types import SimpleNamespace
    from unittest.mock import AsyncMock
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.adapter.database import get_db
    from app.api.admin.sangfor import router, service
    from app.middleware.auth import require_normal_admin
    from app.policy.permissions import ROLE_PERMISSIONS
    from app.middleware.error_handler import app_exception_handler
    from app.port.exceptions import AppException

    app = FastAPI()
    app.include_router(router, prefix="/admin")
    app.add_exception_handler(AppException, app_exception_handler)
    async def db():
        yield SimpleNamespace()
    app.dependency_overrides[get_db] = db
    result = {"items": [], "total": 0, "page": 1, "page_size": 20, "status_counts": {"approved": 2}}
    mocked = AsyncMock(return_value=result)
    monkeypatch.setattr(service, "refunds", mocked)
    monkeypatch.setattr(service, "refund_batches", AsyncMock(return_value=[{"id": 7, "name": "批次"}]))
    client = TestClient(app)
    url = "/admin/cert-products/sangfor/refunds"
    assert client.get(url).status_code == 401
    assert client.get("/admin/cert-products/sangfor/refund-batches").status_code == 401
    app.dependency_overrides[require_normal_admin] = lambda: SimpleNamespace(id=101, role="course_admin")
    assert client.get(url).status_code == 403
    assert client.get("/admin/cert-products/sangfor/refund-batches").status_code == 403
    mocked.assert_not_awaited()
    # A ledger-only role needs no content permission for its batch selector.
    monkeypatch.setitem(ROLE_PERMISSIONS, "ledger_reader", ["order:list"])
    app.dependency_overrides[require_normal_admin] = lambda: SimpleNamespace(id=101, role="ledger_reader")
    assert client.get("/admin/cert-products/sangfor/batches").status_code == 403
    assert client.get("/admin/cert-products/sangfor/refund-batches").json()["data"] == [{"id": 7, "name": "批次"}]
    response = client.get(url, params={"batch_id": 7, "registration_id": 8, "request_kind": "exception", "group": "abnormal",
        "keyword": "考生", "requested_at_start": "2026-10-01T00:00:00+08:00", "page_size": 20})
    assert response.status_code == 200
    assert response.json()["data"]["status_counts"] == {"approved": 2}
    assert mocked.call_args.kwargs["registration_id"] == 8
    assert mocked.call_args.kwargs["group"] == "abnormal"
    assert mocked.call_args.kwargs["requested_at_start"].utcoffset().total_seconds() == 28800
    for params in ({"page": 0}, {"page_size": 101}, {"page_size": 0}, {"batch_id": 0}, {"registration_id": -1},
                   {"group": "other"}, {"status": "requested"}, {"request_kind": "batch_cancel"},
                   {"requested_at_start": "not-a-date"}):
        assert client.get(url, params=params).status_code == 422, params


@pytest.mark.asyncio
@pytest.mark.parametrize("params", [
    {"page": 0}, {"page_size": 101}, {"batch_id": -1}, {"registration_id": 0},
    {"group": "other"}, {"status": "requested"}, {"request_kind": "invalid"},
    {"requested_at_start": datetime(2026, 10, 1)},
    {"requested_at_start": datetime.fromisoformat("2026-10-02T00:00:00+08:00"),
     "requested_at_end": datetime.fromisoformat("2026-10-01T00:00:00+08:00")},
])
async def test_refund_filters_fail_before_database_access(params):
    from app.services.sangfor import SangforService
    with pytest.raises(BusinessException):
        await SangforService().refunds(**params)


@pytest.mark.asyncio
async def test_admin_create_batch_returns_draft_id_for_publish_retry(monkeypatch):
    from types import SimpleNamespace
    from unittest.mock import AsyncMock
    from app.api.admin.sangfor import create_batch, service, PlanService
    from app.schemas.sangfor import SangforBatchCreate

    monkeypatch.setattr(service, "require_product", AsyncMock())
    create = AsyncMock(return_value=SimpleNamespace(id=17))
    monkeypatch.setattr(PlanService, "create_plan", create)
    result = await create_batch(SangforBatchCreate(product_code="SCSA-T", name="测试批次", price_cents=100), admin=SimpleNamespace(id=1))
    assert result.data == {"id": 17}
    assert create.call_args.kwargs["admin_id"] == 1
