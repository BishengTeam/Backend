"""The real reporting workbook and immutable reporting revision boundaries."""

from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.domain.order.src.index import Order
from app.domain.plan.src.index import Plan
from app.domain.renshe.src.index import RensheApplication, RensheApplicationVersion
from app.domain.user.src.index import UserRealname
from app.models.renshe_project import RensheProjectConfig
from app.port.exceptions import BusinessException, ConflictException
from app.schemas.renshe import RensheReviewCreate
from app.schemas.renshe_project import RensheProjectWrite
from app.services.renshe_application import RensheApplicationService
from app.services.renshe_export import ExportCandidate, build_registration_row
from app.services.renshe_graduated import (
    DEGREE_REVIEW_CHECKS as REVIEW_CHECKS,
    validate_education,
)
from app.services.renshe_project import RensheProjectService
from app.services.renshe_reporting import (
    DEGREE_CONDITION,
    EDUCATIONS,
    ETHNICITIES,
    assert_template,
    public_options,
    report_date,
    resolve_contract,
    validate_reporting,
)
from app.services.renshe_review import RensheReviewService
from tests.unit.test_renshe_graduated import draft_payload, graduated_batch, uploads
from tests.unit.test_renshe_projects import isolated as _isolated

isolated = _isolated


def info(**changes):
    return {
        "occupation_status": "职工",
        "ethnicity": "汉族",
        "political_status": "群众",
        "has_work_experience": True,
        "work_start_date": "2022-09-21",
        "occupation_years": 0,
        **changes,
    }


def test_actual_client_workbook_structure_dictionaries_and_condition():
    original = (
        Path(__file__).resolve().parents[3]
        / "DOCS/docs/甲方/认证材料/人设报名模版/报名信息.xlsx"
    )
    source = Path(__file__).resolve().parents[2] / "docs/renshe/报名信息.xlsx"
    if original.exists():
        assert original.read_bytes() == source.read_bytes()
    workbook = load_workbook(source)
    try:
        assert_template(workbook)
        assert workbook["Sheet1"].max_column == 25
        assert workbook["模板示例"]["K1"].value.startswith("(必填)")
        assert workbook["模板示例"]["L1"].value.startswith("(必填)")
        assert set(ETHNICITIES) == {
            row[0]
            for row in workbook["模板说明"].iter_rows(
                min_row=267, max_row=323, max_col=1, values_only=True
            )
        }
        workbook["Sheet1"]["M1"] = "毕业状态"
        with pytest.raises(ValueError, match="标题"):
            assert_template(workbook)
    finally:
        workbook.close()


@pytest.mark.parametrize(
    "changes",
    [
        {"occupation_status": "已毕业"},
        {"ethnicity": ""},
        {"political_status": ""},
        {"has_work_experience": "false"},
        {"work_start_date": None},
        {"occupation_years": None},
        {"occupation_years": -1},
        {"occupation_years": "0.001"},
        {"work_start_date": "2999-01-01"},
        {"has_work_experience": False},
        {
            "has_work_experience": False,
            "work_start_date": None,
            "enrollment_date": "2019-09-01",
        },
    ],
)
def test_reporting_rejects_missing_or_fabricated_fields(changes):
    with pytest.raises(BusinessException):
        validate_reporting(info(**changes), {})


def test_current_work_date_and_years_cannot_be_changed_by_product_policy():
    result = validate_reporting(
        info(),
        {
            "employment_reporting": {
                "date_rule": "enrollment_only",
                "years_rule": "optional",
                "allow_zero_years": False,
            }
        },
    )
    assert float(result["occupation_years"]) == 0
    assert report_date(result, {"date_rule": "enrollment_only"}) == "2022-9-21"
    with pytest.raises(BusinessException, match="参加工作日期"):
        validate_reporting(info(work_start_date=None, enrollment_date="2019-09-01"), {})
    assert report_date(info(work_start_date="2022-01-02")) == "2022-1-2"


@pytest.mark.asyncio
async def test_fixed_rows_ignore_old_export_mapping_and_keep_actual_data(
    isolated, monkeypatch
):
    _, batch = await graduated_batch(isolated, monkeypatch)
    service = RensheApplicationService(isolated.storage)
    saved = await service.save_draft(
        1,
        draft_payload(
            batch,
            reporting_info=info(
                occupation_years="3.25",
                work_unit="真实单位",
                social_insurance_unit="真实参保单位",
            ),
        ),
    )
    async with isolated.db() as db:
        plan = await db.get(Plan, batch.id)
        realname = await db.scalar(select(UserRealname))
        candidate = ExportCandidate(
            saved.id,
            1,
            0,
            service._realname_snapshot(realname),
            {},
            saved.draft_data,
            (),
        )
        plan.renshe_config = {
            **plan.renshe_config,
            "field_mapping": {"1": "email"},
            "export_template": "renshe_roster_v1",
            "attachment_layout": "by_material",
        }
        values = build_registration_row(candidate, plan)
        assert len(values) == 25 and values[0] == "张三"
        assert values[11:13] == ["真实单位", "职工"]
        assert values[17:19] == ["2022-9-21", 3.25]
        assert values[23] == "真实参保单位" and values[20] == DEGREE_CONDITION
        for code, label in EDUCATIONS.items():
            education = {
                **saved.draft_data["education_info"],
                "education": code,
                "institution_type": "vocational",
            }
            if code in {"primary", "junior_high", "senior_high", "other"}:
                with pytest.raises(BusinessException, match="不属于本期"):
                    validate_education(education, plan.renshe_config)
                continue
            validate_education(education, plan.renshe_config)
            row = build_registration_row(
                ExportCandidate(
                    saved.id,
                    1,
                    0,
                    candidate.realname_snapshot,
                    {},
                    {**saved.draft_data, "education_info": education},
                    (),
                ),
                plan,
            )
            assert row[6] == label
        validate_education(
            {**saved.draft_data["education_info"], "education": "higher_vocational"},
            plan.renshe_config,
        )
        with pytest.raises(BusinessException, match="专科及以上"):
            validate_education(
                {**saved.draft_data["education_info"], "education": "technical_school"},
                plan.renshe_config,
            )


@pytest.mark.asyncio
async def test_project_enable_requires_only_business_fields_and_commitment_file(
    isolated, monkeypatch
):
    project, _ = await graduated_batch(isolated, monkeypatch)
    config = project.config.model_dump(
        exclude={
            "field_mapping",
            "export_template",
            "attachment_layout",
            "reporting_policy",
        }
    )
    updated = await RensheProjectService().save(
        RensheProjectWrite(
            chinese_name=project.chinese_name,
            expected_version=project.version,
            status="available",
            config=config,
        ),
        product_id=project.id,
        admin_id=1,
    )
    assert updated.validation_errors == []
    assert (
        updated.config.field_mapping == {}
        and updated.config.attachment_layout == "flat"
    )
    assert updated.config.reporting_policy == {"contract_version": 2}
    config["commitment_template"] = None
    with pytest.raises(BusinessException, match="承诺书"):
        await RensheProjectService().save(
            RensheProjectWrite(
                chinese_name=project.chinese_name,
                expected_version=updated.version,
                status="available",
                config=config,
            ),
            product_id=project.id,
            admin_id=1,
        )


@pytest.mark.asyncio
async def test_old_paid_snapshot_adopts_new_reporting_without_rewriting_order_or_qualification(
    isolated, monkeypatch
):
    project, batch = await graduated_batch(isolated, monkeypatch)
    legacy = project.config.model_dump()
    legacy.pop("selection_id", None)
    legacy["education_levels"] = ["bachelor"]
    legacy["reporting_policy"].pop("contract_version")
    legacy["reporting_policy"].pop("employment_reporting", None)
    legacy["reporting_policy"]["export_values"] = {
        "candidate_source": "社会人员",
        "student_status": "已毕业",
        "application_condition": "已取得本专业毕业证书",
    }
    old_form = draft_payload(batch).model_dump(
        mode="json", exclude={"reporting_info", "reporting_config_id"}
    )
    service = RensheApplicationService(isolated.storage)
    async with isolated.db() as db:
        legacy_row = RensheProjectConfig(
            product_id=project.id, version=2, config=legacy
        )
        db.add(legacy_row)
        await db.flush()
        plan = await db.get(Plan, batch.id)
        plan.renshe_config_id, plan.renshe_config = legacy_row.id, legacy
        plan.status, plan.apply_end = (
            "registration_closed",
            datetime.now(timezone.utc) - timedelta(days=1),
        )
        application = RensheApplication(
            user_id=1,
            plan_id=batch.id,
            status="pending_initial_review",
            draft_data=old_form,
        )
        db.add(application)
        await db.flush()
        realname = await db.scalar(select(UserRealname))
        old_version = RensheApplicationVersion(
            application_id=application.id,
            project_config_id=legacy_row.id,
            version_no=1,
            submitted_by_user_id=1,
            realname_snapshot=service._realname_snapshot(realname),
            student_snapshot={},
            form_data=old_form,
            submitted_at=datetime.now(timezone.utc),
        )
        db.add(old_version)
        await db.flush()
        application.current_version_id = old_version.id
        order = Order(
            user_id=1,
            plan_id=plan.id,
            application_id=application.id,
            product_type=plan.product_type,
            price=12345,
            status="paid",
            paid_at=datetime.now(timezone.utc),
        )
        db.add(order)
        await db.commit()
        legacy_config_id, old_version_id, application_id, order_id = (
            legacy_row.id,
            old_version.id,
            application.id,
            order.id,
        )
    updated = await RensheProjectService().save(
        RensheProjectWrite(
            chinese_name=project.chinese_name,
            expected_version=2,
            status="available",
            config=project.config.model_dump(),
        ),
        product_id=project.id,
        admin_id=1,
    )
    async with isolated.db() as db:
        plan = await db.get(Plan, batch.id)
        options = public_options(await resolve_contract(db, plan), plan.renshe_config)
        assert options["education_levels"] == ["bachelor"]
        assert options["config_id"] == updated.config_id
    review = RensheReviewService()
    with pytest.raises(ConflictException, match="缺少模板填报信息"):
        await review.review(
            reviewer_id=1,
            application_id=application_id,
            stage="initial",
            data=RensheReviewCreate(
                decision="approved",
                expected_version_id=old_version_id,
                checklist={key: True for key in REVIEW_CHECKS},
            ),
        )
    await review.review(
        reviewer_id=1,
        application_id=application_id,
        stage="initial",
        data=RensheReviewCreate(
            decision="rejected",
            expected_version_id=old_version_id,
            required_changes=["reporting_info.occupation_status"],
            reason="补齐模板填报信息",
            correction_due_at=datetime.now(timezone.utc) + timedelta(days=2),
        ),
    )
    await service.save_draft(
        1, draft_payload(batch, reporting_config_id=updated.config_id)
    )
    await uploads(service, batch)
    result = await service.submit(1, application_id)
    assert result.order_id == order_id and result.payment_required is False
    assert result.version.version_no == 2
    async with isolated.db() as db:
        assert (await db.get(Order, order_id)).price == 12345
        assert await db.scalar(select(func.count()).select_from(Order)) == 1
        old = await db.get(RensheApplicationVersion, old_version_id)
        assert old.form_data == old_form and old.project_config_id == legacy_config_id
        new = await db.get(RensheApplicationVersion, result.version.id)
        assert new.project_config_id == legacy_config_id
        assert (
            new.form_data["commitment_template"] == updated.config.commitment_template
        )
        assert (await db.get(Plan, batch.id)).renshe_config == legacy
    await review.review(
        reviewer_id=1,
        application_id=application_id,
        stage="initial",
        data=RensheReviewCreate(
            decision="approved",
            expected_version_id=result.version.id,
            checklist={key: True for key in REVIEW_CHECKS},
        ),
    )
