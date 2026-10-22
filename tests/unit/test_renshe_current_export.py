"""Real ZIPs, current code rules and safe handling of queued/finished tasks."""

import io
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.domain.order.src.index import Order
from app.domain.renshe.src.index import RensheExportVolume
from app.port.config import settings
from app.port.exceptions import ConflictException
from app.schemas.renshe import RensheReviewCreate
from app.schemas.renshe_project import RensheProjectWrite
from app.services.order_fulfillment import OrderFulfillmentService
from app.services.renshe_application import RensheApplicationService
from app.services.renshe_export import RensheExportService
from app.services.renshe_export_definition import (
    export_instructions,
    material_archive_path,
)
from app.services.renshe_graduated import DEGREE_REVIEW_CHECKS as REVIEW_CHECKS, save_template
from app.services.renshe_project import RensheProjectService
from app.services.renshe_project import graduated_configuration
from app.services.renshe_review import RensheReviewService
from tests.unit.test_renshe_graduated import graduated_batch, draft_payload, uploads
from tests.unit.test_renshe_projects import isolated as _isolated

isolated = _isolated


async def approved(env, monkeypatch):
    project, batch = await graduated_batch(env, monkeypatch)
    service = RensheApplicationService(env.storage)
    await uploads(service, batch)
    draft = await service.save_draft(1, draft_payload(batch))
    result = await service.submit(1, draft.id)
    async with env.db() as db:
        order = await db.get(Order, result.order_id)
        order.payment_channel = "wechat"
        order.offline_status = None
        order.status, order.transaction_id, order.paid_at = (
            "paid",
            "current-export-test",
            datetime.now(timezone.utc),
        )
        await OrderFulfillmentService().on_paid(db, order)
        await db.commit()
    await RensheReviewService().review(
        reviewer_id=1,
        application_id=draft.id,
        stage="initial",
        data=RensheReviewCreate(
            decision="approved",
            expected_version_id=result.version.id,
            checklist={key: True for key in REVIEW_CHECKS},
        ),
    )
    return project, batch, result


@pytest.mark.asyncio
async def test_actual_zip_matches_shared_preview_and_keeps_real_files(
    isolated, monkeypatch
):
    _, batch, _ = await approved(isolated, monkeypatch)
    exporter = RensheExportService(isolated.storage)
    job = await exporter.create_job(plan_id=batch.id, admin_id=1)
    _, candidates = await exporter._load_candidates(job.id)
    assert await exporter.process_next_job()
    done = await exporter.get_job(job.id)
    assert done.status == "succeeded", done.last_error
    async with isolated.db() as db:
        volume = await db.scalar(
            select(RensheExportVolume).where(RensheExportVolume.job_id == job.id)
        )
        path = isolated.storage._local_path(volume.storage_key)
    candidate = candidates[0]
    identity = candidate.realname_snapshot["id_card_number"]
    expected = {"报名信息.xlsx", "工作经历.xlsx"} | {
        material_archive_path(identity, item.kind, item.extension)
        for item in candidate.materials
    }
    with zipfile.ZipFile(path) as archive:
        assert set(archive.namelist()) == expected
        assert (
            archive.read("工作经历.xlsx")
            == (Path(settings.RENSHE_TEMPLATE_DIR) / "工作经历.xlsx").read_bytes()
        )
        assert archive.read(f"{identity}_学历证明.jpeg").startswith(
            b"\xff\xd8\xff"
        )
        assert archive.read(f"{identity}_二寸免冠照片.jpeg").startswith(
            b"\xff\xd8\xff"
        )
        book = load_workbook(io.BytesIO(archive.read("报名信息.xlsx")))
        assert book.sheetnames == ["Sheet1", "模板示例", "模板说明", "申报条件"]
        assert book["Sheet1"].max_column == 25
        assert (
            book["Sheet1"]["R2"].value == "2022-9-21"
            and book["Sheet1"]["S2"].value == 0
        )
        book.close()
    instructions = export_instructions()
    assert len(instructions["columns"]) == 25
    assert all(
        item.rsplit("/", 1)[-1] in instructions["file_tree"]
        for item in instructions["example_paths"]
    )
    assert "原样" in instructions["work_description"]


@pytest.mark.asyncio
async def test_rule_upgrade_rejects_queued_task_but_preserves_completed_download(
    isolated, monkeypatch
):
    _, batch, _ = await approved(isolated, monkeypatch)
    exporter = RensheExportService(isolated.storage)
    first = await exporter.create_job(plan_id=batch.id, admin_id=1)
    await exporter.process_next_job()
    async with isolated.db() as db:
        volume = await db.scalar(
            select(RensheExportVolume).where(RensheExportVolume.job_id == first.id)
        )
        volume_id, key = volume.id, volume.storage_key
    original_bytes = isolated.storage._local_path(key).read_bytes()
    queued = await exporter.create_job(plan_id=batch.id, admin_id=1)
    monkeypatch.setattr(
        "app.services.renshe_export_definition.RULE_ID", "test-upgrade-current-rules"
    )
    with pytest.raises(ConflictException, match="旧规则在途任务"):
        await exporter.create_job(plan_id=batch.id, admin_id=1)
    await exporter.process_next_job()
    failed = await exporter.get_job(queued.id)
    assert failed.status == "failed" and "新建导出任务" in failed.last_error
    with pytest.raises(ConflictException, match="新建导出任务"):
        await exporter.retry_job(job_id=queued.id, admin_id=1)
    latest = await exporter.create_job(plan_id=batch.id, admin_id=1)
    await exporter.process_next_job()
    assert (await exporter.get_job(latest.id)).status == "succeeded"
    assert (await exporter.get_job(first.id)).status == "succeeded"
    assert isolated.storage._local_path(key).read_bytes() == original_bytes
    signed = AsyncMock(return_value="https://example.test/private-export")
    monkeypatch.setattr(isolated.storage, "signed_get_url", signed)
    assert (
        await exporter.volume_signed_url(
            volume_id=volume_id, admin_id=1, ip_address=None
        )
    ).url == "https://example.test/private-export"
    assert signed.call_args.args[0] == key


@pytest.mark.asyncio
async def test_selected_project_metadata_save_preserves_rules_and_commitment(
    isolated, monkeypatch
):
    selected = graduated_configuration()
    project = await RensheProjectService().save(
        RensheProjectWrite(chinese_name="信息安全中级", status="draft",
                           occupation_code=selected["occupation_code"], skill_level=selected["skill_level"]),
        admin_id=1,
    )
    monkeypatch.setattr(
        settings, "RENSHE_TEMPLATE_DIR", str(isolated.root / "templates")
    )
    config = graduated_configuration()
    template = save_template("确认承诺书.pdf", b"%PDF-1.4\nfixture")
    config["commitment_template"] = template
    updated = await RensheProjectService().save(
        RensheProjectWrite(
            chinese_name=project.chinese_name,
            status="available",
            expected_version=project.version,
            config=config,
        ),
        product_id=project.id,
        admin_id=1,
    )
    assert updated.config.commitment_template == template
    expected = graduated_configuration()
    assert updated.config.qualification_rule == expected["qualification_rule"]
    assert updated.config.education_levels == expected["education_levels"]
    assert [item.model_dump() for item in updated.config.materials] == expected["materials"]
    minimal = await RensheProjectService().save(
        RensheProjectWrite(
            chinese_name=project.chinese_name,
            status="available",
            expected_version=updated.version,
        ),
        product_id=project.id,
        admin_id=1,
    )
    assert minimal.config.commitment_template == template
