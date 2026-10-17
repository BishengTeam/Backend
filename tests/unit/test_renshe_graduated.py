"""Graduated-degree lifecycle with a private local store and isolated SQL DB."""

from datetime import datetime, timedelta, timezone
from pathlib import Path
import shutil
import zipfile

import pytest
from openpyxl import load_workbook
from sqlalchemy import delete, select

from tests.unit.test_renshe_projects import isolated as _isolated, verified_student, create_batch
from app.domain.order.src.index import Order
from app.domain.plan.src.index import Plan
from app.domain.user.src.index import UserRealname, UserStudent
from app.domain.renshe.src.index import (
    RensheApplication,
    RensheApplicationVersion,
    RensheExportVolume,
    RensheMaterial,
)
from app.port.config import settings
from app.port.exceptions import (
    BusinessException,
    ConflictException,
    ValidationException,
)
from app.schemas.renshe import RensheDraftUpsert, RensheReviewCreate
from app.schemas.renshe_project import RensheProjectWrite
from app.services.renshe_application import RensheApplicationService
from app.services.renshe_graduated import (
    DEGREE_REVIEW_CHECKS as REVIEW_CHECKS,
    save_template,
    activation_errors,
    china_today,
)
from app.services.renshe_project import (
    RensheProjectService,
    graduated_configuration,
    default_configuration,
)
from app.services.renshe_review import RensheReviewService
from app.services.renshe_export import RensheExportService
from app.services.plan import PlanService
from app.services.order_fulfillment import OrderFulfillmentService
from app.services.renshe_reporting import default_policy, DEGREE_CONDITION, assert_template

isolated = _isolated


async def graduated_batch(env, monkeypatch):
    root = env.root / "templates"
    root.mkdir()
    for name in ("报名信息.xlsx", "工作经历.xlsx"):
        shutil.copyfile(Path(settings.RENSHE_TEMPLATE_DIR) / name, root / name)
    monkeypatch.setattr(settings, "RENSHE_TEMPLATE_DIR", str(root))
    config = graduated_configuration()
    config["reporting_policy"] = {
        **default_policy(),
        "version": "fixture-v2",
        "confirmed_by_client": True,
        "commitment_template": save_template("测试模板.pdf", b"%PDF-1.4\nfixture"),
        "employment_reporting": {"confirmed": True, "date_rule": "work_or_enrollment", "years_rule": "required", "allow_zero_years": True},
    }
    project = await RensheProjectService().save(
        RensheProjectWrite(
            chinese_name="已毕业学历型", status="available", config=config
        ),
        admin_id=1,
    )
    batch = await create_batch(project)
    await PlanService().publish_plan(batch.id, project.code, admin_id=1)
    await verified_student(env)
    async with env.db() as db:
        await db.execute(delete(UserStudent))
        identity = await db.scalar(select(UserRealname))
        identity.user_type = "social"
        await db.commit()
    return project, batch


def draft_payload(batch, **changes):
    values = dict(
        plan_id=batch.id,
        contact_phone="13800138000",
        email="graduate@example.com",
        mailing_address="毕业报名地址",
        reporting_config_id=batch.renshe_config_id,
        reporting_info={"occupation_status": "失业人员", "ethnicity": "汉族", "political_status": "群众", "has_work_experience": True, "work_start_date": "2022-09-21", "occupation_years": 0},
        education_info={
            "graduated": True,
            "school": "测试大学",
            "institution_type": "general_higher",
            "education": "bachelor",
            "major": "网络空间安全",
            "graduation_date": "2023-06-30",
            "certificate_number": "TEST-20230001",
        },
    )
    values.update(changes)
    return RensheDraftUpsert(**values)


async def uploads(service, batch, *, exclude=()):
    for kind in [item["kind"] for item in batch.renshe_config["materials"]]:
        if kind in {"id_card_front", "id_card_back", "portrait", *exclude}:
            continue
        pdf = kind in {"education_report", "xuexin_registration", "education_report", "commitment_letter"}
        await service.upload_application_material(
            1,
            batch.id,
            kind,
            filename="实际原文件.pdf" if pdf else "证明材料.jpeg",
            content_type="application/pdf" if pdf else "image/jpeg",
            content=b"%PDF-1.4\nfixture" if pdf else b"\xff\xd8\xfffixture",
        )


@pytest.mark.asyncio
async def test_graduate_without_student_identity_full_revision_review_and_export(
    isolated, monkeypatch
):  # noqa: F811
    env = isolated
    project, batch = await graduated_batch(env, monkeypatch)
    service = RensheApplicationService(env.storage)
    await uploads(service, batch)
    draft = await service.save_draft(1, draft_payload(batch))
    result = await service.submit(1, draft.id)
    async with env.db() as db:
        order = await db.get(Order, result.order_id)
        order.status, order.transaction_id, order.paid_at = (
            "paid",
            "graduate-test",
            datetime.now(timezone.utc),
        )
        await OrderFulfillmentService().on_paid(db, order)
        await db.commit()
        version = await db.get(RensheApplicationVersion, result.version.id)
        assert (
            version.realname_snapshot["user_type"] == "social"
            and version.student_snapshot == {}
        )
        assert (
            version.form_data["education_info"]["certificate_number"] == "TEST-20230001"
        )
    review = RensheReviewService()
    with pytest.raises(BusinessException, match="七项"):
        await review.review(
            reviewer_id=1,
            application_id=draft.id,
            stage="initial",
            data=RensheReviewCreate(
                expected_version_id=result.version.id, decision="approved"
            ),
        )
    with pytest.raises(BusinessException, match="截止时间"):
        await review.review(
            reviewer_id=1,
            application_id=draft.id,
            stage="initial",
            data=RensheReviewCreate(
                expected_version_id=result.version.id,
                decision="rejected",
                reason="专业信息需补正",
                required_changes=["education_info.major"],
            ),
        )
    deadline = datetime.now(timezone.utc) + timedelta(days=3)
    await review.review(
        reviewer_id=1,
        application_id=draft.id,
        stage="initial",
        data=RensheReviewCreate(
            expected_version_id=result.version.id,
            decision="rejected",
            reason="专业按毕业证原文填写",
            required_changes=["education_info.major"],
            correction_due_at=deadline,
        ),
    )
    async with env.db() as db:
        plan = await db.get(Plan, batch.id)
        plan.status, plan.apply_end = (
            "registration_closed",
            datetime.now(timezone.utc) - timedelta(days=1),
        )
        await db.commit()
    updated = draft_payload(batch)
    updated.education_info.major = "信息安全"
    await service.save_draft(1, updated)
    await uploads(
        service,
        batch,
        exclude=("commitment_letter",),
    )
    resubmitted = await service.submit(1, draft.id)
    assert not resubmitted.payment_required and resubmitted.order_id == result.order_id
    assert resubmitted.version.version_no == 2
    async with env.db() as db:
        old = await db.get(RensheApplicationVersion, result.version.id)
        assert old.form_data["education_info"]["major"] == "网络空间安全"
        assert (await db.get(RensheApplication, draft.id)).correction_due_at is None
        old_material = await db.scalar(
            select(RensheMaterial).where(RensheMaterial.version_id == old.id)
        )
        assert env.storage._local_path(old_material.storage_key).is_file()
    checks = {key: True for key in REVIEW_CHECKS}
    with pytest.raises(ConflictException, match="版本已变化"):
        await review.review(
            reviewer_id=1,
            application_id=draft.id,
            stage="initial",
            data=RensheReviewCreate(
                expected_version_id=result.version.id,
                decision="approved",
                checklist=checks,
            ),
        )
    approved = await review.review(
        reviewer_id=1,
        application_id=draft.id,
        stage="initial",
        data=RensheReviewCreate(
            expected_version_id=resubmitted.version.id,
            decision="approved",
            checklist=checks,
        ),
    )
    assert approved.checklist == checks
    with pytest.raises(ConflictException, match="尚未成功导出"):
        await review.review(
            reviewer_id=1,
            application_id=draft.id,
            stage="external",
            data=RensheReviewCreate(
                expected_version_id=resubmitted.version.id, decision="approved"
            ),
        )
    exporter = RensheExportService(env.storage)
    job = await exporter.create_job(plan_id=batch.id, admin_id=1)
    await exporter.process_next_job()
    done = await exporter.get_job(job.id)
    assert done.status == "succeeded", done.last_error
    assert done.items[0].version_id == resubmitted.version.id
    async with env.db() as db:
        volume = await db.scalar(
            select(RensheExportVolume).where(RensheExportVolume.job_id == job.id)
        )
        path = env.storage._local_path(volume.storage_key)
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        assert any(name.endswith("_承诺书.pdf") for name in names)
        assert any(name.endswith("_二寸免冠照片.jpeg") for name in names)
        assert any(name.endswith("_学信网电子注册表.pdf") for name in names)
        assert any(name.endswith("_学历证明.jpeg") for name in names)
        assert all("/" not in name for name in names)
        assert not any("学生证" in name for name in names)
        target = env.root / "graduated.xlsx"
        target.write_bytes(archive.read("报名信息.xlsx"))
    workbook = load_workbook(target)
    assert workbook["Sheet1"].cell(2, 6).value == "社会人员"
    assert workbook["Sheet1"].cell(2, 13).value == "失业人员"
    assert workbook["Sheet1"].cell(2, 18).value == "2022-9-21"
    assert workbook["Sheet1"].cell(2, 19).value == 0
    assert workbook["Sheet1"].cell(2, 21).value == DEGREE_CONDITION
    assert_template(workbook)
    workbook.close()
    await review.review(
        reviewer_id=1,
        application_id=draft.id,
        stage="external",
        data=RensheReviewCreate(
            expected_version_id=resubmitted.version.id, decision="approved"
        ),
    )


@pytest.mark.asyncio
async def test_missing_material_size_scope_identity_and_expired_corrections(
    isolated, monkeypatch
):  # noqa: F811
    env = isolated
    _, batch = await graduated_batch(env, monkeypatch)
    service = RensheApplicationService(env.storage)
    draft = await service.save_draft(1, draft_payload(batch))
    await uploads(service, batch, exclude=("commitment_letter",))
    with pytest.raises(BusinessException, match="承诺书"):
        await service.submit(1, draft.id)
    with pytest.raises(ValidationException, match="大小"):
        await service.upload_application_material(
            1,
            batch.id,
            "education_report",
            filename="large.pdf",
            content_type="application/pdf",
            content=b"%PDF-1.4" + b"a" * (20 * 1024 * 1024),
        )
    with pytest.raises(ValueError):
        draft_payload(
            batch,
            education_info={
                **draft_payload(batch).education_info.model_dump(mode="json"),
                "graduated": False,
            },
        )
    with pytest.raises(ValueError):
        draft_payload(
            batch,
            education_info={
                **draft_payload(batch).education_info.model_dump(mode="json"),
                "graduation_date": (china_today() + timedelta(days=1)).isoformat(),
            },
        )
    with pytest.raises(ValueError):
        draft_payload(batch, real_name="冒用他人")
    with pytest.raises(ValueError):
        draft_payload(
            batch, materials={"commitment_letter": {"storage_key": "other/user/file"}}
        )
    async with env.db() as db:
        application = await db.get(RensheApplication, draft.id)
        application.status = "initial_rejected"
        application.correction_due_at = datetime.now(timezone.utc) - timedelta(
            seconds=1
        )
        await db.commit()
    with pytest.raises(ConflictException, match="补正期限"):
        await service.save_draft(1, draft_payload(batch))
    with pytest.raises(ConflictException, match="补正期限"):
        await uploads(service, batch)
    with pytest.raises(ConflictException, match="补正期限"):
        await service.submit(1, draft.id)


def test_policy_activation_and_templates_never_assume_client_approval(
    tmp_path, monkeypatch
):
    monkeypatch.setattr(settings, "RENSHE_TEMPLATE_DIR", str(tmp_path))
    assert activation_errors(default_configuration())
    assert activation_errors(graduated_configuration())
    with pytest.raises(BusinessException):
        save_template("template.pdf", b"wrong-format")
    with pytest.raises(BusinessException):
        save_template("template.jpg", b"\xff\xd8\xfffixture")
    metadata = save_template("../template.pdf", b"%PDF-1.4\nfixture")
    assert (
        metadata["filename"].startswith("commitment_")
        and "/" not in metadata["filename"]
    )
    assert (tmp_path / metadata["filename"]).is_file()


@pytest.mark.asyncio
async def test_new_rules_reject_unadapted_combinations_and_changed_uploads(
    isolated, monkeypatch
):  # noqa: F811
    env = isolated
    project, batch = await graduated_batch(env, monkeypatch)
    config = project.config.model_dump()
    config["skill_level"] = "高级工"
    with pytest.raises(BusinessException, match="四级"):
        await RensheProjectService().save(
            RensheProjectWrite(
                chinese_name="未适配三级", status="available", config=config
            ),
            admin_id=1,
        )
    service = RensheApplicationService(env.storage)
    await uploads(service, batch)
    draft = await service.save_draft(1, draft_payload(batch))
    metadata = draft.draft_data["materials"]["education_report"]
    env.storage._local_path(metadata["storage_key"]).write_bytes(b"%PDF-1.4\nmodified")
    with pytest.raises(ConflictException, match="内容已变化"):
        await service.submit(1, draft.id)
    async with env.db() as db:
        assert not (await db.execute(select(RensheApplicationVersion))).scalars().all()
        assert not (await db.execute(select(Order))).scalars().all()


@pytest.mark.asyncio
async def test_graduated_http_upload_download_and_draft_ownership(
    isolated, monkeypatch
):  # noqa: F811
    from httpx import ASGITransport, AsyncClient
    from types import SimpleNamespace
    from app.main import app
    from app.middleware.auth import get_current_user

    _, batch = await graduated_batch(isolated, monkeypatch)
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(id=1)
    try:
        async with AsyncClient(
            transport=ASGITransport(app=app), base_url="http://test"
        ) as client:
            options = await client.get(f"/api/renshe/batches/{batch.id}/reporting-options")
            assert options.status_code == 200, options.text
            assert options.json()["data"]["config_id"] == batch.renshe_config_id
            assert options.json()["data"]["application_condition"] == DEGREE_CONDITION
            response = await client.post(
                f"/api/renshe/batches/{batch.id}/materials/commitment_letter",
                files={"file": ("承诺书.pdf", b"%PDF-1.4\nfixture", "application/pdf")},
            )
            assert response.status_code == 200, response.text
            uploaded = response.json()["data"]
            assert uploaded["original_filename"] == "承诺书.pdf"
            downloaded = await client.get(
                f"/api/renshe/batches/{batch.id}/commitment-template"
            )
            assert downloaded.status_code == 200 and downloaded.content.startswith(
                b"%PDF-"
            )
            body = draft_payload(batch).model_dump(mode="json")
            response = await client.post(
                "/api/renshe/applications/draft",
                json={
                    **body,
                    "materials": {
                        "commitment_letter": {"storage_key": "other/user/file"}
                    },
                },
            )
            assert response.status_code == 422
            response = await client.post("/api/renshe/applications/draft", json=body)
            assert response.status_code == 200, response.text
            assert (
                response.json()["data"]["draft_data"]["materials"]["commitment_letter"][
                    "storage_key"
                ]
                == uploaded["storage_key"]
            )
            app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(id=2)
            assert (
                await client.get(
                    f"/api/renshe/applications/{uploaded['application_id']}"
                )
            ).status_code == 404
    finally:
        app.dependency_overrides.pop(get_current_user, None)
