"""Each adapted trade must survive real admission, review and workbook export."""

import hashlib
import io
import zipfile
from copy import deepcopy
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.domain.order.src.index import Order
from app.domain.plan.src.index import Plan
from app.domain.renshe.src.index import RensheExportVolume
from app.port.config import settings
from app.port.exceptions import BusinessException
from app.schemas.renshe import RensheReviewCreate
from app.schemas.renshe_project import RensheProjectWrite
from app.services.order_fulfillment import OrderFulfillmentService
from app.services.plan import PlanService
from app.services.renshe_application import RensheApplicationService
from app.services.renshe_catalog import (
    ADAPTED_DEGREE_OCCUPATIONS,
    combinations,
    find_selection,
)
from app.services.renshe_export import RensheExportService
from app.services.renshe_export_definition import prepare_reporting_workbook
from app.services.renshe_graduated import DEGREE_REVIEW_CHECKS, validate_education
from app.services.renshe_project import (
    RensheProjectService,
    configuration_errors,
    selection_configuration,
)
from app.services.renshe_reporting import public_options
from app.services.renshe_review import RensheReviewService
from tests.unit.test_renshe_graduated import draft_payload, graduated_batch, uploads
from tests.unit.test_renshe_projects import create_batch
from tests.unit.test_renshe_projects import isolated as _isolated

isolated = _isolated


@pytest.mark.asyncio
@pytest.mark.parametrize("occupation", sorted(ADAPTED_DEGREE_OCCUPATIONS))
async def test_each_degree_trade_uses_its_own_rules_through_real_zip(isolated, monkeypatch, occupation):
    original, _ = await graduated_batch(isolated, monkeypatch)
    chosen = find_selection(f"{occupation}:4:学历型")
    service = RensheProjectService()
    project = await service.save(RensheProjectWrite(
        chinese_name="自定义名称不能决定职业", selection_id=chosen["selection_id"], status="available",
        commitment_template=original.config.commitment_template,
    ), admin_id=1)
    assert project.config.occupation_code == chosen["occupation_code"]
    assert project.config.occupation_name == chosen["occupation_name"]
    assert project.admission_requirements["rule_profile"]["related_majors"] == chosen["related_majors"]
    batch = await create_batch(project)
    await PlanService().publish_plan(batch.id, project.code, admin_id=1)
    config = deepcopy(batch.renshe_config)
    options = public_options({"config_id": project.config_id, "config": {"reporting_policy": {"commitment_template": project.config.commitment_template}}}, config)
    assert options["institution_types"] == chosen["institution_types"]
    assert options["rule_profile"]["occupation_name"] == chosen["occupation_name"]
    applicant = RensheApplicationService(isolated.storage)
    await uploads(applicant, batch)
    payload = draft_payload(batch)
    if occupation == "ai_trainer":
        payload.education_info.institution_type = "vocational"
        payload.education_info.education = "secondary_vocational"
        payload.education_info.major = "电子技术应用"
    draft = await applicant.save_draft(1, payload)
    submitted = await applicant.submit(1, draft.id)
    async with isolated.db() as db:
        order = await db.get(Order, submitted.order_id)
        order.payment_channel = "wechat"
        order.offline_status = None
        order.status, order.paid_at, order.transaction_id = "paid", datetime.now(UTC), f"multi-degree-{occupation}"
        await OrderFulfillmentService().on_paid(db, order)
        await db.commit()
    await RensheReviewService().review(reviewer_id=1, application_id=draft.id, stage="initial", data=RensheReviewCreate(
        decision="approved", expected_version_id=submitted.version.id, checklist={key: True for key in DEGREE_REVIEW_CHECKS},
    ))
    exporter = RensheExportService(isolated.storage)
    job = await exporter.create_job(plan_id=batch.id, admin_id=1)
    await exporter.process_next_job()
    done = await exporter.get_job(job.id)
    assert done.status == "succeeded", done.last_error
    async with isolated.db() as db:
        volume = await db.scalar(select(RensheExportVolume).where(RensheExportVolume.job_id == job.id))
        path = isolated.storage._local_path(volume.storage_key)
    with zipfile.ZipFile(path) as archive:
        workbook = load_workbook(io.BytesIO(archive.read("报名信息.xlsx")))
        sheet = workbook["Sheet1"]
        assert sheet["N2"].value == chosen["occupation_name"]
        assert sheet["O2"].value == "四级"
        assert sheet["T2"].value == chosen["occupation_code"]
        assert sheet["U2"].value == config["rule_profile"]["source_conditions"][0]
        assert sheet["V2"].value == "学历型"
        assert sheet["R2"].value == "2022-9-21" and sheet["S2"].value == 0
        assert sheet.max_column == 25
        workbook.close()
    # A display edit must not alter the published batch or its rules.
    updated = await service.save(RensheProjectWrite(chinese_name="修改展示信息", status="available", expected_version=project.version), product_id=project.id, admin_id=1)
    assert updated.config.model_dump() == project.config.model_dump()
    async with isolated.db() as db:
        assert (await db.get(Plan, batch.id)).renshe_config == config


def test_ai_degree_never_implicitly_accepts_unconfirmed_university_scope():
    config = selection_configuration(find_selection("ai_trainer:4:学历型"))
    info = draft_payload(type("Batch", (), {"id": 1, "renshe_config_id": 1})()).education_info.model_dump(mode="json")
    with pytest.raises(BusinessException, match="院校类型"):
        validate_education(info, config)
    forged = deepcopy(config)
    forged["institution_types"].append("general_higher")
    assert configuration_errors(forged)


@pytest.mark.asyncio
async def test_previously_pending_draft_can_select_newly_adapted_rules(isolated, monkeypatch):
    original, _ = await graduated_batch(isolated, monkeypatch)
    chosen = find_selection("human_resources:4:学历型")
    old = selection_configuration({**chosen, "adapted": False})
    old["commitment_template"] = original.config.commitment_template
    service = RensheProjectService()
    draft = await service.save(RensheProjectWrite(chinese_name="上线前草稿", config=old), admin_id=1)
    assert draft.config.qualification_rule == "unadapted"
    ready = await service.save(RensheProjectWrite(chinese_name=draft.chinese_name, selection_id=chosen["selection_id"], status="available", expected_version=draft.version), product_id=draft.id, admin_id=1)
    assert ready.config.qualification_rule != "unadapted"
    assert ready.config.occupation_name == "企业人力资源管理师"
    assert ready.config.commitment_template == original.config.commitment_template


def test_reporting_dictionary_contains_all_current_trades_without_changing_source():
    source = Path(__file__).resolve().parents[2] / "docs/renshe/报名信息.xlsx"
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source)
    try:
        prepare_reporting_workbook(workbook)
        prepare_reporting_workbook(workbook)  # retries must not duplicate conditions
        assert workbook.sheetnames == ["Sheet1", "模板示例", "模板说明", "申报条件"]
        definition = workbook.defined_names["renshe_current_occupations"]
        sheet_name, cells = next(definition.destinations)
        codes = {cell.value for row in workbook[sheet_name][cells] for cell in row}
        adapted = [item for item in combinations() if item["adapted"]]
        assert codes == {item["occupation_code"] for item in adapted}
        rows = list(workbook["申报条件"].iter_rows(min_row=2, max_col=2, values_only=True))
        for item in adapted:
            pair = (item["occupation_code"], item["source_conditions"][0])
            assert rows.count(pair) == 1
            indexes = [i for i, row in enumerate(rows) if row[0] == item["occupation_code"]]
            assert indexes == list(range(min(indexes), max(indexes) + 1))
        assert any("T2" in rule.sqref and rule.formula1 == "renshe_current_occupations" for rule in workbook["Sheet1"].data_validations.dataValidation)
        assert any("U2" in rule.sqref and "OFFSET" in (rule.formula1 or "") for rule in workbook["Sheet1"].data_validations.dataValidation)
    finally:
        workbook.close()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before


@pytest.mark.asyncio
async def test_current_template_download_uses_the_same_dictionary_as_exports(monkeypatch):
    from app.api.admin.renshe import download_current_registration_template
    monkeypatch.setattr(settings, "RENSHE_TEMPLATE_DIR", str(Path(__file__).resolve().parents[2] / "docs/renshe"))
    response = await download_current_registration_template(_admin=None)
    assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    workbook = load_workbook(io.BytesIO(response.body))
    try:
        assert "renshe_current_occupations" in workbook.defined_names
        assert workbook["Sheet1"].max_column == 25
        assert "filename*=UTF-8''" in response.headers["content-disposition"]
    finally:
        workbook.close()
