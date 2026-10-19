"""Grade-three branch boundaries and real per-path material/export evidence."""

import io
import zipfile
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.domain.renshe.src.index import RensheExportVolume, RensheMaterial
from app.port.exceptions import BusinessException
from app.schemas.renshe import RensheDraftUpsert, RensheReviewCreate
from app.schemas.renshe_project import RensheProjectWrite
from app.services.plan import PlanService
from app.services.renshe_application import RensheApplicationService
from app.services.renshe_catalog import ADAPTED_DEGREE_OCCUPATIONS, find_selection
from app.services.renshe_employment import anniversary
from app.services.renshe_export import RensheExportService
from app.services.renshe_graduated import china_today, validate_admission, review_checks
from app.services.renshe_level_three import PATHS, material_requirements, selected_path
from app.services.renshe_project import RensheProjectService, selection_configuration
from app.services.renshe_review import RensheReviewService
from tests.unit.test_renshe_employment import pay
from tests.unit.test_renshe_graduated import graduated_batch
from tests.unit.test_renshe_projects import create_batch, isolated as _isolated

isolated = _isolated


def test_all_six_conditions_and_title_grade_match_original_client_workbook():
    from pathlib import Path
    from app.services.renshe_catalog import CONDITIONS
    source = Path(__file__).resolve().parents[2] / "docs/renshe/报名信息.xlsx"
    workbook = load_workbook(source)
    try:
        conditions = [row[1] for row in workbook["申报条件"].iter_rows(values_only=True) if row[0] == "4-04-04-02-网络与信息安全管理员-信息安全管理员-高级工"]
        assert len(conditions) == 6
        assert {value for values in CONDITIONS[3].values() for value in values} == set(conditions)
        assert any("Q2" in rule.sqref and "初级职称" in (rule.formula1 or "") for rule in workbook["Sheet1"].data_validations.dataValidation)
    finally:
        workbook.close()


def config(path, occupation="information_security"):
    return selection_configuration(find_selection(f"{occupation}:3:{PATHS[path]['type']}"))


def payload(path, plan_id=1):
    today = china_today()
    branch = PATHS[path]
    years = branch["years"]
    start = anniversary(today, -(years + 1)).isoformat()
    info = {"admission_path": path, "cultural_education": "associate", "work_experiences": []}
    if branch.get("work"):
        info["work_experiences"] = [{"work_unit": "样例单位", "position": "实际岗位", "occupation_name": "信息安全管理员", "start_date": start, "end_date": None}]
    if branch.get("skill"):
        info["skill_certificate"] = {"occupation_name": "信息安全管理员", "grade": 4, "certificate_number": "TEST-L4", "issued_on": start}
    if branch.get("title"):
        info["title_certificate"] = {"title_name": "助理工程师", "level": "primary", "major": "信息安全", "certificate_number": "TEST-TITLE", "issued_on": start}
    data = {"plan_id": plan_id, "contact_phone": "13800138000", "email": "three@example.test", "mailing_address": "实际地址",
            "qualification_info": info,
            "reporting_info": {"occupation_status": "职工", "ethnicity": "汉族", "political_status": "群众", "has_work_experience": True,
                               "work_start_date": start, "occupation_years": years, "work_unit": "样例单位", "social_insurance_unit": "样例单位"}}
    if branch.get("education"):
        data["education_info"] = {"graduated": True, "school": "样例院校", "institution_type": "technical" if path == "technical_senior" else "vocational",
                                  "education": "technical_school" if path == "technical_senior" else "higher_vocational", "major": "信息安全",
                                  "graduation_date": "2020-07-01", "certificate_number": "TEST-DEGREE"}
        if path == "technical_senior":
            data["education_info"]["training_level"] = "senior_worker"
        if path == "evaluated_college":
            data["education_info"]["evaluation_reference"] = "备案等级认定评价机构，初审核验原始依据"
    return data


@pytest.mark.parametrize("occupation", sorted(ADAPTED_DEGREE_OCCUPATIONS))
@pytest.mark.parametrize("path", list(PATHS))
def test_each_trade_and_branch_has_own_rules(occupation, path):
    policy = config(path, occupation)
    result = validate_admission(payload(path), policy)
    assert result["eligibility_snapshot"]["required_years"] == PATHS[path]["years"]
    assert result["eligibility_snapshot"]["application_condition"] == selected_path(policy, payload(path))["condition"]
    required = {i["kind"] for i in material_requirements(policy, payload(path))}
    assert ("work_proof" in required) == bool(PATHS[path].get("work"))
    assert ("professional_title_certificate" in required) == bool(PATHS[path].get("title"))
    assert ("graduation_certificate" in required) == bool(PATHS[path].get("education"))


@pytest.mark.parametrize("path", ["work_10", "skill_4_work_4", "title_work_1"])
def test_year_boundaries_overlap_and_certificate_date_are_enforced(path):
    data = payload(path)
    branch = PATHS[path]
    threshold = anniversary(china_today(), -branch["years"])
    info = data["qualification_info"]
    info["work_experiences"][0]["start_date"] = (threshold + timedelta(days=2)).isoformat()
    info["work_experiences"] *= 2
    with pytest.raises(BusinessException, match="累计相关工作"):
        validate_admission(data, config(path))
    data = payload(path)
    info = data["qualification_info"]
    if branch.get("skill") or branch.get("title"):
        certificate = info["skill_certificate"] if branch.get("skill") else info["title_certificate"]
        certificate["issued_on"] = china_today().isoformat()
        with pytest.raises(BusinessException, match="证书取得后"):
            validate_admission(data, config(path))
    else:
        data["reporting_info"]["occupation_years"] = 5
        with pytest.raises(BusinessException, match="职业年限"):
            validate_admission(data, config(path))


def test_degree_layers_and_branch_switching_do_not_borrow_other_rules():
    for path in ["technical_senior", "evaluated_college", "skill_4_degree"]:
        data = payload(path)
        data["education_info"]["graduated"] = False
        with pytest.raises(BusinessException, match="毕业证"):
            validate_admission(data, config(path))
    data = payload("technical_senior")
    data["education_info"].pop("training_level")
    with pytest.raises(BusinessException, match="高级工班"):
        validate_admission(data, config("technical_senior"))
    data = payload("evaluated_college")
    data["education_info"].pop("evaluation_reference")
    with pytest.raises(BusinessException, match="评估论证"):
        validate_admission(data, config("evaluated_college"))
    data = payload("skill_4_degree")
    data["qualification_info"]["skill_certificate"]["grade"] = 5
    with pytest.raises(BusinessException, match="四级"):
        validate_admission(data, config("skill_4_degree"))
    data = payload("skill_4_work_4")
    data["qualification_info"]["admission_path"] = "skill_4_degree"
    with pytest.raises(BusinessException, match="毕业证"):
        validate_admission(data, config("skill_4_degree"))


@pytest.mark.asyncio
@pytest.mark.parametrize("path", list(PATHS))
async def test_full_path_material_review_and_actual_zip(isolated, monkeypatch, path):
    original, _ = await graduated_batch(isolated, monkeypatch)
    service = RensheProjectService()
    project = await service.save(RensheProjectWrite(chinese_name=f"三级-{path}", selection_id=config(path)["selection_id"], commitment_template=original.config.commitment_template, status="available"), admin_id=1)
    batch = await create_batch(project)
    await PlanService().publish_plan(batch.id, project.code, admin_id=1)
    applicant = RensheApplicationService(isolated.storage)
    data = payload(path, batch.id)
    if PATHS[path]["type"] == "学历型":
        data["reporting_info"].update(has_work_experience=False, work_start_date=None, occupation_years=0, work_unit=None, social_insurance_unit=None)
    draft = await applicant.save_draft(1, RensheDraftUpsert.model_validate(data))
    requirements = material_requirements(batch.renshe_config, data)
    for requirement in requirements:
        kind = requirement["kind"]
        if kind.startswith("id_card_"):
            continue
        with pytest.raises(BusinessException, match="请上传"):
            await applicant.submit(1, draft.id)
        pdf = ".pdf" in requirement["extensions"] and kind not in {"registration_portrait", "graduation_certificate"}
        await applicant.upload_application_material(1, batch.id, kind, filename="原始.pdf" if pdf else "原始.jpg", content_type="application/pdf" if pdf else "image/jpeg", content=b"%PDF-1.4\nfixture" if pdf else b"\xff\xd8\xfffixture")
    submitted = await applicant.submit(1, draft.id)
    await pay(isolated, submitted)
    reviewer = RensheReviewService()
    with pytest.raises(BusinessException, match="逐项"):
        await reviewer.review(reviewer_id=1, application_id=draft.id, stage="initial", data=RensheReviewCreate(decision="approved", expected_version_id=submitted.version.id))
    await reviewer.review(reviewer_id=1, application_id=draft.id, stage="initial", data=RensheReviewCreate(decision="approved", expected_version_id=submitted.version.id, checklist={key: True for key in review_checks(batch.renshe_config)}))
    exporter = RensheExportService(isolated.storage)
    job = await exporter.create_job(plan_id=batch.id, admin_id=1)
    await exporter.process_next_job()
    finished = await exporter.get_job(job.id)
    assert finished.status == "succeeded", finished.last_error
    async with isolated.db() as db:
        volume = await db.scalar(select(RensheExportVolume).where(RensheExportVolume.job_id == job.id))
        archive_path = isolated.storage._local_path(volume.storage_key)
    with zipfile.ZipFile(archive_path) as archive:
        workbook = load_workbook(io.BytesIO(archive.read("报名信息.xlsx")))
        row = workbook["Sheet1"]
        if PATHS[path]["type"] == "学历型":
            assert row["R2"].value is None
            assert row["S2"].value == 0
        assert row["O2"].value == "三级"
        assert row["U2"].value == selected_path(batch.renshe_config, data)["condition"]
        assert row["V2"].value == PATHS[path]["type"]
        assert row["P2"].value == ("信息安全" if path == "title_work_1" else None)
        assert row["Q2"].value == ("初级职称" if path == "title_work_1" else None)
        workbook.close()
        files = archive.namelist()
        assert any("_工龄证明" in item for item in files) == bool(PATHS[path].get("work"))
        assert any("_职称证书" in item for item in files) == bool(PATHS[path].get("title"))
        assert any("_学历证明" in item for item in files) == bool(PATHS[path].get("education"))
    await reviewer.review(reviewer_id=1, application_id=draft.id, stage="external", data=RensheReviewCreate(decision="approved", expected_version_id=submitted.version.id))


@pytest.mark.asyncio
async def test_paid_branch_correction_reuses_order_and_only_snapshots_current_materials(isolated, monkeypatch):
    original, _ = await graduated_batch(isolated, monkeypatch)
    project = await RensheProjectService().save(RensheProjectWrite(chinese_name="三级复合型补正", selection_id=config("skill_4_work_4")["selection_id"], commitment_template=original.config.commitment_template, status="available"), admin_id=1)
    batch = await create_batch(project)
    await PlanService().publish_plan(batch.id, project.code, admin_id=1)
    applicant = RensheApplicationService(isolated.storage)
    data = payload("skill_4_work_4", batch.id)
    draft = await applicant.save_draft(1, RensheDraftUpsert.model_validate(data))

    async def upload(kind, extensions):
        pdf = ".pdf" in extensions and kind not in {"registration_portrait", "graduation_certificate"}
        await applicant.upload_application_material(1, batch.id, kind, filename="材料.pdf" if pdf else "材料.jpg", content_type="application/pdf" if pdf else "image/jpeg", content=b"%PDF-1.4\nfixture" if pdf else b"\xff\xd8\xfffixture")

    for item in material_requirements(batch.renshe_config, data):
        if not item["kind"].startswith("id_card_"):
            await upload(item["kind"], item["extensions"])
    submitted = await applicant.submit(1, draft.id)
    await pay(isolated, submitted)
    await RensheReviewService().review(reviewer_id=1, application_id=draft.id, stage="initial", data=RensheReviewCreate(
        decision="rejected", expected_version_id=submitted.version.id, reason="请核实实际申报路径", required_changes=["qualification_info.admission_path", "education_info.school"], correction_due_at=datetime.now(timezone.utc) + timedelta(days=3)))
    data = payload("skill_4_degree", batch.id)
    await applicant.save_draft(1, RensheDraftUpsert.model_validate(data))
    for item in material_requirements(batch.renshe_config, data):
        if item["kind"] in {"graduation_certificate", "education_report"}:
            await upload(item["kind"], item["extensions"])
    corrected = await applicant.submit(1, draft.id)
    assert corrected.order_id == submitted.order_id and not corrected.payment_required
    assert corrected.version.id != submitted.version.id
    async with isolated.db() as db:
        rows = (await db.execute(select(RensheMaterial).where(RensheMaterial.application_id == draft.id))).scalars().all()
    assert {r.kind for r in rows if r.version_id == corrected.version.id} == set(selected_path(batch.renshe_config, data)["material_kinds"])
    assert "work_proof" in {r.kind for r in rows if r.version_id == submitted.version.id}
