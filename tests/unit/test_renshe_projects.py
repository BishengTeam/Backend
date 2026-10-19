"""Isolated SQL/storage evidence for multi-project admission; no external I/O."""

from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock
import zipfile

import pytest
from openpyxl import load_workbook
from sqlalchemy import CheckConstraint, MetaData, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.adapter.database import Base
from app.domain.order.src.index import Order
from app.domain.plan.src.index import Plan
from app.domain.renshe.src.index import (
    RensheExportVolume,
)
from app.domain.user.src.index import AdminUser, User, UserRealname, UserStudent
from app.integrations.renshe_storage import RensheObjectStorage
from app.integrations.wechat_pay import WechatPayRefund
from app.models.renshe_project import RensheProjectConfig
from app.port.config import settings
from app.port.exceptions import BusinessException, ConflictException, NotFoundException
from app.schemas.plan import PlanCreate, PlanUpdate
from app.schemas.renshe import RensheDraftUpsert, RensheReviewCreate, RensheRefundCreate
from app.schemas.renshe_project import RensheProjectWrite
from app.services.plan import PlanService
from app.services.renshe_application import RensheApplicationService
from app.services.renshe_export import RensheExportService
from app.services.renshe_project import (
    RensheProjectService,
    batch_configuration,
    configuration_errors,
    default_configuration,
)
from app.services.renshe_review import RensheReviewService
from app.services.renshe_refund import RensheRefundService
from app.services.order_fulfillment import OrderFulfillmentService


@pytest.fixture
async def isolated(monkeypatch, tmp_path):
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'projects.sqlite'}")
    names = {
        "cert_product_catalog",
        "cert_product",
        "renshe_project_config",
        "user",
        "user_realname",
        "user_student",
        "admin_user",
        "plan",
        "order",
        "cert_consultation_group",
    }
    names.update(name for name in Base.metadata.tables if name.startswith("renshe_"))
    metadata = MetaData()
    for table in Base.metadata.tables.values():
        table.to_metadata(metadata)
    admin_table = metadata.tables["admin_user"]
    for constraint in list(admin_table.constraints):
        if constraint.name == "ck_admin_user_role":
            admin_table.constraints.remove(constraint)
    admin_table.append_constraint(
        CheckConstraint(
            "role IN ('super_admin', 'quiz_admin', 'cert_admin', 'course_admin', 'teacher')"
        )
    )
    async with engine.begin() as connection:
        await connection.run_sync(
            lambda db: metadata.create_all(
                db, tables=[metadata.tables[name] for name in names]
            )
        )
    factory = async_sessionmaker(engine, expire_on_commit=False)

    @asynccontextmanager
    async def context():
        async with factory() as session:
            yield session

    monkeypatch.setattr("app.adapter.database.get_db_ctx", context)

    for module in (
        "renshe_project",
        "plan",
        "renshe_application",
        "renshe_export",
        "renshe_review",
        "renshe_refund",
        "renshe_cleanup",
        "order",
        "admin_cert_product",
    ):
        monkeypatch.setattr(f"app.services.{module}.get_db_ctx", context)
    # These tests exercise persisted, pre-rollout project policies. Seed their
    # old availability without weakening the production publication gate.
    from app.services.renshe_graduated import activation_errors, is_graduated
    monkeypatch.setattr("app.services.renshe_project.activation_errors", lambda config: activation_errors(config) if is_graduated(config) else [])
    monkeypatch.setattr("app.services.renshe_application.ensure_accepted", AsyncMock())
    monkeypatch.setattr(
        "app.services.h3c_registration.H3cRegistrationService.on_order_paid",
        AsyncMock(return_value=False),
    )
    monkeypatch.setattr(
        "app.services.nisp_registration.NispRegistrationService.on_order_paid",
        AsyncMock(return_value=False),
    )
    monkeypatch.setattr(
        PlanService, "_reject_dedicated_certification_batch", AsyncMock()
    )
    monkeypatch.setattr(settings, "UPLOAD_DIR", str(tmp_path / "uploads"))
    monkeypatch.setattr(settings, "RENSHE_COMMITMENT_TEMPLATE_DIR", "")
    monkeypatch.setattr(settings, "RENSHE_STORAGE_TYPE", "local")
    template_dir = Path(__file__).parents[2] / "docs" / "renshe"
    monkeypatch.setattr(settings, "RENSHE_TEMPLATE_DIR", str(template_dir))
    async with context() as db:
        db.add(
            AdminUser(
                id=1,
                username="multi-project-test",
                display_name="测试管理员",
                password_hash="test",
                role="super_admin",
            )
        )
        db.add(User(id=1, openid="multi-project-student", is_active=True))
        await db.commit()
    try:
        yield SimpleNamespace(db=context, storage=RensheObjectStorage(), root=tmp_path)
    finally:
        await engine.dispose()


def project_data(name="信息安全中级", **changes):
    config = default_configuration()
    config.update(changes)
    return RensheProjectWrite(chinese_name=name, status="available", config=config)


async def create_batch(project, name="第一期", **changes):
    now = datetime.now(timezone.utc)
    values = dict(
        name=name,
        project_config_id=project.config_id,
        apply_start=now - timedelta(hours=1),
        apply_end=now + timedelta(days=10),
        exam_date=now + timedelta(days=15),
        price_cents=50100,
        exam_location="测试考点",
        description="测试报名说明",
        contact_name="老师",
        contact_phone="13800138000",
    )
    values.update(changes)
    return await PlanService().create_plan(
        project.code, PlanCreate(**values), admin_id=1
    )


async def publish_historical_batch(env, batch):
    """Represent a batch published before the new admission rule existed."""
    async with env.db() as db:
        plan = await db.get(Plan, batch.id)
        plan.status = "published"
        plan.published_at = datetime.now(timezone.utc)
        await db.commit()


async def verified_student(env):
    sources = {}
    for requirement in default_configuration()["materials"]:
        pdf = requirement["kind"] == "xuexin_registration"
        material = await env.storage.save_source(
            user_id=1,
            kind=requirement["kind"],
            filename="proof.pdf" if pdf else "proof.jpg",
            content_type="application/pdf" if pdf else "image/jpeg",
            data=b"%PDF-1.4\nfixture" if pdf else b"\xff\xd8\xfffixture",
        )
        sources[requirement["kind"]] = material.storage_key
    async with env.db() as db:
        db.add(
            UserRealname(
                user_id=1,
                real_name="张三",
                user_type="student",
                id_card_number="11010519491231002X",
                gender="男",
                birth_date="1949-12-31",
                ethnicity="汉族",
                political_status="群众",
                status="verified",
                id_card_front_oss=sources["id_card_front"],
                id_card_back_oss=sources["id_card_back"],
                avatar_oss=sources["portrait"],
            )
        )
        db.add(
            UserStudent(
                user_id=1,
                education="bachelor",
                school="测试学校",
                major="网络安全",
                enrollment_date=date(2023, 9, 1),
                status="verified",
                student_card_oss=sources["student_card"],
                enrollment_pdf_oss=sources["xuexin_registration"],
                degree_cert_oss=sources["education_proof"],
            )
        )
        await db.commit()


async def enroll(env, batch):
    from app.domain.renshe.src.index import RensheApplication
    # A persisted pre-rollout application may continue under its frozen rule;
    # creating a new legacy application after rollout is deliberately blocked.
    async with env.db() as db:
        db.add(RensheApplication(user_id=1, plan_id=batch.id, status="draft"))
        await db.commit()
    service = RensheApplicationService(env.storage)
    draft = await service.save_draft(
        1,
        RensheDraftUpsert(
            plan_id=batch.id,
            contact_phone="13800138000",
            mailing_address="测试地址",
            email="test@example.com",
            reporting_info={"occupation_status": "职工", "ethnicity": "汉族", "political_status": "群众", "has_work_experience": True, "work_start_date": "2022-09-21", "occupation_years": 0},
        ),
    )
    submission = await service.submit(1, draft.id)
    async with env.db() as db:
        order = await db.get(Order, submission.order_id)
        assert order.product_type == batch.product_type
        order.status, order.transaction_id, order.paid_at = (
            "paid",
            f"transaction-{order.id}",
            datetime.now(timezone.utc),
        )
        await OrderFulfillmentService().on_paid(db, order)
        await db.commit()
    await RensheReviewService().review(
        reviewer_id=1,
        application_id=draft.id,
        stage="initial",
        data=RensheReviewCreate(decision="approved"),
    )
    return submission


@pytest.mark.asyncio
async def test_two_projects_complete_separate_admission_exports_external_review_and_refund(
    isolated,
):
    env = isolated
    await verified_student(env)
    service = RensheProjectService()
    first = await service.save(project_data(), admin_id=1)
    second = await service.save(
        project_data(
            "信息安全第二项目",
            export_template="renshe_roster_v1",
            attachment_layout="by_material",
        ),
        admin_id=1,
    )
    assert first.id != second.id and first.code != second.code
    # Each project has multiple independent batches.
    for project in (first, second):
        await create_batch(project, "第二期")
    batches = [await create_batch(project) for project in (first, second)]
    for batch in batches:
        await publish_historical_batch(env, batch)
        with pytest.raises(ConflictException, match="本期新报名"):
            await RensheApplicationService(env.storage).save_draft(1, RensheDraftUpsert(
                plan_id=batch.id, contact_phone="13800138000", email="test@example.com", mailing_address="地址"
            ))
    submissions = [await enroll(env, batch) for batch in batches]
    from app.services.admin_cert_product import AdminCertProductService

    stats = next(
        item
        for item in await AdminCertProductService().get_stats()
        if item.type == "renshe"
    )
    assert (stats.product_count, stats.active_batch_count, stats.total_enrolled) == (
        2,
        2,
        2,
    )
    listed = await RensheApplicationService(env.storage).list_applications(1)
    assert {item.plan.id for item in listed.items} == {batch.id for batch in batches}
    exports = RensheExportService(env.storage)
    for index, (batch, submitted) in enumerate(zip(batches, submissions)):
        job = await exports.create_job(
            plan_id=batch.id, admin_id=1, application_ids=[submitted.application.id]
        )
        assert await exports.process_next_job()
        result = await exports.get_job(job.id)
        assert result.status == "succeeded", result.last_error
        async with env.db() as db:
            volume = await db.scalar(
                select(RensheExportVolume).where(RensheExportVolume.job_id == job.id)
            )
            path = env.storage._local_path(volume.storage_key)
        with zipfile.ZipFile(path) as archive:
            target = env.root / f"registration-{index}.xlsx"
            target.write_bytes(archive.read("报名信息.xlsx"))
            workbook = load_workbook(target)
            values = [cell.value for cell in workbook["Sheet1"][2]]
            if index == 0:
                assert values[13:15] == ["信息安全管理员", "四级"]
                assert any(
                    name.startswith("110105") for name in archive.namelist()
                )
            else:
                assert values[13:15] == ["信息安全管理员", "四级"]
                assert any(
                    name.startswith("110105")
                    for name in archive.namelist()
                )
            workbook.close()
        await RensheReviewService().review(
            reviewer_id=1,
            application_id=submitted.application.id,
            stage="external",
            data=RensheReviewCreate(decision="approved"),
        )
    # Refunds still use the dedicated full-refund flow and original order price.
    submitted = submissions[1]
    refunds = RensheRefundService(wechat_pay=SimpleNamespace(mch_id="test-merchant"))
    requested = await refunds.request_refund(
        user_id=1,
        application_id=submitted.application.id,
        data=RensheRefundCreate(
            request_kind="exception",
            reason_code="test",
            reason_detail="测试环境例外退款",
        ),
    )
    prepared = await refunds._prepare_submission(
        refund_id=requested.id, actor_type="admin", actor_id=1, system_only=False
    )
    provider = WechatPayRefund(
        out_trade_no=prepared.out_trade_no,
        transaction_id=prepared.transaction_id,
        out_refund_no=prepared.out_refund_no,
        refund_id="test-refund",
        status="SUCCESS",
        amount_total=50100,
        amount_refund=50100,
        currency="CNY",
        mchid="test-merchant",
        success_time=datetime.now(timezone.utc),
    )
    applied = await refunds._apply_provider_result(
        provider, source="query", expected_refund_id=requested.id, allow_success=True
    )
    assert applied.refund.status == "succeeded"
    async with env.db() as db:
        assert (await db.get(Order, submitted.order_id)).status == "refunded"
        assert (await db.get(Order, submissions[0].order_id)).status == "paid"


@pytest.mark.asyncio
async def test_version_freeze_explicit_draft_upgrade_disable_and_immutable_rows(
    isolated,
):
    service = RensheProjectService()
    first = await service.save(project_data(), admin_id=1)
    published = await create_batch(first)
    draft = await create_batch(first, "草稿")
    await publish_historical_batch(isolated, published)
    config = first.config.model_dump()
    config["education_levels"] = ["bachelor"]
    second = await service.save(
        RensheProjectWrite(
            chinese_name="新展示名称",
            status="available",
            config=config,
            expected_version=1,
        ),
        product_id=first.id,
        admin_id=1,
    )
    assert second.version == 2
    plans = await PlanService().list_plans(first.code)
    assert all(item.renshe_config_id == first.config_id for item in plans)
    assert all(item.project_name == first.chinese_name for item in plans)
    updated = await PlanService().update_plan(
        draft.id, PlanUpdate(project_config_id=second.config_id), first.code, admin_id=1
    )
    assert updated.renshe_config["education_levels"] == ["bachelor"]
    for payload in (
        PlanUpdate(project_config_id=second.config_id),
        PlanUpdate(exam_type="补考"),
        PlanUpdate(project_code=first.code),
    ):
        with pytest.raises(BusinessException):
            await PlanService().update_plan(published.id, payload, first.code)
    disabled = await service.save(
        RensheProjectWrite(
            chinese_name=second.chinese_name,
            status="disabled",
            config=second.config,
            expected_version=2,
        ),
        product_id=first.id,
        admin_id=1,
    )
    with pytest.raises(BusinessException):
        await create_batch(disabled, "禁止创建")
    with pytest.raises(BusinessException):
        await PlanService().publish_plan(draft.id, first.code)
    assert any(
        item.id == published.id
        for item in await PlanService().list_renshe_plans(public=True)
    )
    # Existing published admission remains usable after project disable.
    await verified_student(isolated)
    await enroll(isolated, published)
    async with isolated.db() as db:
        frozen = await db.get(RensheProjectConfig, first.config_id)
        frozen.config = config
        with pytest.raises(ValueError, match="不可修改"):
            await db.commit()


@pytest.mark.asyncio
async def test_draft_capability_gate_cross_project_forgery_and_deletion(isolated):
    service = RensheProjectService()
    data = project_data()
    data.config.education_levels = []
    data.status = "draft"
    draft = await service.save(data, admin_id=1)
    assert draft.validation_errors
    with pytest.raises(BusinessException):
        await create_batch(draft)
    data.status = "available"
    data.expected_version = draft.version
    with pytest.raises(BusinessException):
        await service.save(data, product_id=draft.id, admin_id=1)
    await service.delete(draft.id, admin_id=1)
    first, second = [
        await service.save(project_data(name), admin_id=1)
        for name in ("项目一", "项目二")
    ]
    with pytest.raises(BusinessException, match="归属不一致"):
        await create_batch(first, project_config_id=second.config_id)
    batch = await create_batch(first)
    with pytest.raises(NotFoundException):
        await PlanService().update_plan(batch.id, PlanUpdate(name="伪造"), second.code)
    data = RensheProjectWrite(
        chinese_name=first.chinese_name,
        config=first.config,
        status="draft",
        expected_version=first.version,
    )
    await service.save(data, product_id=first.id, admin_id=1)
    with pytest.raises(BusinessException):
        await service.delete(first.id, admin_id=1)
    async with isolated.db() as db:
        original = await db.get(Plan, batch.id)
        original.renshe_config_id = second.config_id
        await db.commit()
    with pytest.raises(ConflictException):
        await PlanService().publish_plan(batch.id, first.code)


@pytest.mark.asyncio
async def test_export_keeps_requested_material_version_after_resubmission_and_rejects_other_batch(
    isolated,
):
    await verified_student(isolated)
    project = await RensheProjectService().save(project_data(), admin_id=1)
    batch = await create_batch(project)
    await publish_historical_batch(isolated, batch)
    submission = await enroll(isolated, batch)
    exports = RensheExportService(isolated.storage)
    other = await create_batch(project, "其他批次")
    await publish_historical_batch(isolated, other)
    with pytest.raises(BusinessException):
        await exports.create_job(
            plan_id=other.id, admin_id=1, application_ids=[submission.application.id]
        )
    job = await exports.create_job(plan_id=batch.id, admin_id=1)
    await RensheReviewService().return_for_correction(
        reviewer_id=1,
        application_id=submission.application.id,
        data=RensheReviewCreate(
            decision="rejected", reason="测试补正", required_changes=["email"]
        ),
    )
    service = RensheApplicationService(isolated.storage)
    await service.save_draft(
        1,
        RensheDraftUpsert(
            plan_id=batch.id,
            contact_phone="13800138000",
            mailing_address="新地址",
            email="new@example.com",
            reporting_info={"occupation_status": "职工", "ethnicity": "汉族", "political_status": "群众", "has_work_experience": True, "work_start_date": "2022-09-21", "occupation_years": 0},
        ),
    )
    newer = await service.submit(1, submission.application.id)
    assert newer.version.id != submission.version.id
    assert not newer.payment_required and newer.order_id == submission.order_id
    await exports.process_next_job()
    result = await exports.get_job(job.id)
    assert result.status == "succeeded", result.last_error
    assert result.items[0].version_id == submission.version.id
    async with isolated.db() as db:
        volume = await db.scalar(
            select(RensheExportVolume).where(RensheExportVolume.job_id == job.id)
        )
        with zipfile.ZipFile(
            isolated.storage._local_path(volume.storage_key)
        ) as archive:
            target = isolated.root / "old-version.xlsx"
            target.write_bytes(archive.read("报名信息.xlsx"))
    workbook = load_workbook(target)
    assert workbook["Sheet1"].cell(2, 25).value == "test@example.com"
    workbook.close()
    await RensheReviewService().review(
        reviewer_id=1,
        application_id=submission.application.id,
        stage="initial",
        data=RensheReviewCreate(decision="approved"),
    )
    with pytest.raises(ConflictException, match="尚未成功导出"):
        await RensheReviewService().review(
            reviewer_id=1,
            application_id=submission.application.id,
            stage="external",
            data=RensheReviewCreate(decision="approved"),
        )


@pytest.mark.asyncio
async def test_generated_project_cannot_use_generic_order_or_wrong_education(isolated):
    from app.schemas.order import OrderCreate
    from app.services.order import OrderService

    project = await RensheProjectService().save(
        project_data(education_levels=["master"]), admin_id=1
    )
    with pytest.raises(BusinessException, match="只能通过人社报名"):
        await OrderService().create_order(
            1, OrderCreate(order_kind="certification", product_type=project.code)
        )
    batch = await create_batch(project)
    await publish_historical_batch(isolated, batch)
    await verified_student(isolated)
    from app.domain.renshe.src.index import RensheApplication
    async with isolated.db() as db:
        db.add(RensheApplication(user_id=1, plan_id=batch.id, status="draft"))
        await db.commit()
    with pytest.raises(BusinessException, match="学历不符合"):
        await RensheApplicationService(isolated.storage).save_draft(
            1,
            RensheDraftUpsert(
                plan_id=batch.id,
                contact_phone="13800138000",
                mailing_address="地址",
                email="test@example.com",
            ),
        )


@pytest.mark.parametrize(
    "change",
    [
        {"qualification_rule": "arbitrary-script"},
        {"exam_types": ["未知类型"]},
        {"application_types": ["工作经历型"]},
        {"form_fields": ["email"]},
        {"materials": []},
        {"education_levels": []},
    ],
)
def test_unadapted_capabilities_are_never_available(change):
    config = default_configuration()
    config.update(change)
    assert configuration_errors(config)


def test_missing_historical_properties_never_use_default_export_values():
    with pytest.raises(ConflictException):
        batch_configuration(Plan(product_type="RS-ZY", name="未核实历史批次"))


@pytest.mark.asyncio
async def test_audit_is_read_only_and_separates_unconfirmed_history(
    isolated, monkeypatch
):
    from scripts import audit_renshe_projects
    from sqlalchemy import func

    monkeypatch.setattr(audit_renshe_projects, "get_db_ctx", isolated.db)
    project = await RensheProjectService().save(project_data(), admin_id=1)
    batch = await create_batch(project)
    async with isolated.db() as db:
        uncertain = Plan(
            product_type="RS-ZY",
            name="需人工核实",
            occupation_name="未知职业",
            status="draft",
        )
        db.add(uncertain)
        await db.commit()
        count_before = await db.scalar(select(func.count()).select_from(Plan))
    report = await audit_renshe_projects.audit()
    assert report["versioned_plan_ids"] == [batch.id]
    assert report["manual_reconciliation_plan_ids"] == [uncertain.id]
    async with isolated.db() as db:
        assert await db.scalar(select(func.count()).select_from(Plan)) == count_before
