"""试题导入模板（甲方 xlsx）与新题型单元测试。"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook

from app.domain.community.src.rule.quiz import (
    QuizRuleViolation,
    answer_score_ratio,
    answers_match,
    normalize_question_payload,
    normalize_question_score,
    normalize_question_tags,
)
from app.services.admin_quiz import AdminQuizService


TEMPLATE_HEADER = [
    "题目（必填）", "题型（必填）", "选项A（必填）", "选项B（必填）",
    "选项C", "选项D", "选项E", "选项F", "选项G", "选项H", "正确答案（必填）",
    "解析", "分数（必填）", "分类（必填）", "标签",
]


def _build_workbook(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(TEMPLATE_HEADER)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _parse(rows: list[list[object]]):
    return AdminQuizService()._parse_import_rows_xlsx(_build_workbook(rows))


# ── 领域规则：新题型 ──


def test_uncertain_choice_allows_single_answer_in_draft_and_publish() -> None:
    normalized = normalize_question_payload(
        question_type="uncertain_choice",
        question_text="下列哪些正确？",
        options={"A": "1", "B": "2", "C": "3"},
        correct_answer=["A"],
        explanation="解析",
        require_publishable=True,
    )
    assert normalized.correct_answer == ["A"]


def test_uncertain_choice_supports_five_options() -> None:
    normalized = normalize_question_payload(
        question_type="uncertain_choice",
        question_text="五选项",
        options={"A": "1", "B": "2", "C": "3", "D": "4", "E": "5"},
        correct_answer=["A", "B", "C", "D", "E"],
    )
    assert normalized.correct_answer == ["A", "B", "C", "D", "E"]


def test_multiple_choice_still_requires_two_answers_on_publish() -> None:
    with pytest.raises(QuizRuleViolation, match="至少 2 个答案"):
        normalize_question_payload(
            question_type="multiple_choice",
            question_text="多选",
            options={"A": "1", "B": "2", "C": "3"},
            correct_answer=["A"],
            require_publishable=True,
        )


def test_matching_normalize_and_exact_match() -> None:
    normalized = normalize_question_payload(
        question_type="matching",
        question_text="配对",
        options={"A": "红豆生南国", "B": "海内存知己"},
        correct_answer=["春来发几枝", "天涯若比邻"],
    )
    assert normalized.correct_answer == ["春来发几枝", "天涯若比邻"]
    assert answers_match(
        "matching",
        ["春来发几枝", "天涯若比邻"],
        normalized.correct_answer,
        options=normalized.options,
    )
    assert not answers_match(
        "matching",
        ["天涯若比邻", "春来发几枝"],
        normalized.correct_answer,
        options=normalized.options,
    )
    assert (
        answer_score_ratio(
            "matching",
            ["天涯若比邻", "春来发几枝"],
            normalized.correct_answer,
            options=normalized.options,
        )
        == 0
    )


def test_matching_rejects_count_mismatch_and_duplicates() -> None:
    with pytest.raises(QuizRuleViolation, match="数量必须与左侧选项数量一致"):
        normalize_question_payload(
            question_type="matching",
            question_text="配对",
            options={"A": "左1", "B": "左2"},
            correct_answer=["右1"],
        )
    with pytest.raises(QuizRuleViolation, match="不能重复"):
        normalize_question_payload(
            question_type="matching",
            question_text="配对",
            options={"A": "左1", "B": "左2"},
            correct_answer=["右1", "右1"],
        )


def test_ordering_requires_permutation_of_options() -> None:
    normalized = normalize_question_payload(
        question_type="ordering",
        question_text="排序",
        options={"A": "中秋", "B": "端午", "C": "重阳"},
        correct_answer=["端午", "中秋", "重阳"],
    )
    assert answers_match(
        "ordering", ["端午", "中秋", "重阳"], normalized.correct_answer,
        options=normalized.options,
    )
    assert not answers_match(
        "ordering", ["中秋", "端午", "重阳"], normalized.correct_answer,
        options=normalized.options,
    )
    with pytest.raises(QuizRuleViolation, match="一种排列"):
        normalize_question_payload(
            question_type="ordering",
            question_text="排序",
            options={"A": "中秋", "B": "端午", "C": "重阳"},
            correct_answer=["端午", "端午", "重阳"],
        )


def test_option_keys_accept_up_to_h_but_not_i() -> None:
    normalized = normalize_question_payload(
        question_type="single_choice",
        question_text="八选项",
        options={
            "A": "1", "B": "2", "C": "3", "D": "4",
            "E": "5", "F": "6", "G": "7", "H": "8",
        },
        correct_answer="H",
    )
    assert normalized.correct_answer == "H"
    with pytest.raises(QuizRuleViolation, match="A~H"):
        normalize_question_payload(
            question_type="single_choice",
            question_text="九选项",
            options={
                "A": "1", "B": "2", "C": "3", "D": "4",
                "E": "5", "F": "6", "G": "7", "H": "8", "I": "9",
            },
            correct_answer="A",
        )


def test_score_and_tags_normalization() -> None:
    assert normalize_question_score("8") == pytest.approx(8.0)
    assert normalize_question_score(None) is None
    assert normalize_question_tags("文学，简单、常识") == ["文学", "简单", "常识"]
    with pytest.raises(QuizRuleViolation, match="0 至 100"):
        normalize_question_score("120")
    with pytest.raises(QuizRuleViolation, match="最多"):
        normalize_question_tags(["a", "b", "c", "d", "e", "f", "g", "h", "i", "j", "k"])


def test_normalize_payload_carries_score_and_tags() -> None:
    normalized = normalize_question_payload(
        question_type="single_choice",
        question_text="题干",
        options={"A": "1", "B": "2", "C": "3"},
        correct_answer="A",
        score="12.5",
        tags="文学，简单",
    )
    assert float(normalized.score) == 12.5
    assert normalized.tags == ["文学", "简单"]


# ── xlsx 解析：甲方模板 ──


def test_parse_full_template_all_eight_types() -> None:
    rows, errors = _parse([
        ["食醋是什么味道的？", "单选题", "酸", "甜", "苦", "辣", None, None, None, None, "A", None, 8, "趣味知识测验", "简单"],
        ["一年中有几个季节？", "多选题", "春", "夏", "秋", "冬", None, None, None, None, "ABCD", None, 8, "趣味知识测验", "简单"],
        ["下列哪些对？", "不定项选择题", "1", "2", "3", "4", "5", "6", "7", None, "ABCDEFG", None, 6, "测试", None],
        ["荔枝多吃上火？", "判断题", "对", "错", None, None, None, None, None, None, "A", None, 8, "常识", None],
        ["马铃薯是果实？", "判断题", "√", "×", None, None, None, None, None, None, "错", "马铃薯是蔬菜", 8, "常识", "较难"],
        ["一年有_____个月，有_____个季节？", "填空题", "12", "4|四", None, None, None, None, None, None, None, None, 6, "常识", None],
        ["中国的情人节是哪一天？", "问答题", None, None, None, None, None, None, None, None, "七夕", None, 12, "常识", "必知必会"],
        ["请把下列的诗句匹配对应", "配对题", "红豆生南国  ->天涯若比邻", "海内存知己  ->更上一层楼", "床前明月光  ->春来发几枝", "欲穷千里目  ->疑是地上霜", None, None, None, None, "CADB", None, 12, "常识", None],
        ["请根据时间的先后，一一对应排序", "排序题", "中秋节", "端午节", "重阳节", "元宵节", None, None, None, None, "DBAC", None, 12, "常识", None],
    ])
    assert errors == []
    assert len(rows) == 9
    parsed = {item.question_text: item for _, item in rows}

    single = parsed["食醋是什么味道的？"]
    assert single.question_type.value == "single_choice"
    assert single.correct_answer == "A"
    assert single.score == 8
    assert single.tags == ["简单"]

    multiple = parsed["一年中有几个季节？"]
    assert multiple.question_type.value == "multiple_choice"
    assert multiple.correct_answer == ["A", "B", "C", "D"]

    uncertain = parsed["下列哪些对？"]
    assert uncertain.question_type.value == "uncertain_choice"
    assert uncertain.correct_answer == ["A", "B", "C", "D", "E", "F", "G"]
    assert set(uncertain.options) == {"A", "B", "C", "D", "E", "F", "G"}

    judge_letter = parsed["荔枝多吃上火？"]
    assert judge_letter.question_type.value == "judge"
    assert judge_letter.correct_answer == "A"
    judge_text = parsed["马铃薯是果实？"]
    assert judge_text.correct_answer == "B"

    fill = parsed["一年有_____个月，有_____个季节？"]
    assert fill.question_type.value == "fill_blank"
    assert fill.correct_answer == [["12"], ["4", "四"]]

    essay = parsed["中国的情人节是哪一天？"]
    assert essay.question_type.value == "essay"
    assert essay.correct_answer == "七夕"

    matching = parsed["请把下列的诗句匹配对应"]
    assert matching.question_type.value == "matching"
    assert matching.options == {
        "A": "红豆生南国", "B": "海内存知己", "C": "床前明月光", "D": "欲穷千里目",
    }
    assert matching.correct_answer == [
        "春来发几枝", "天涯若比邻", "疑是地上霜", "更上一层楼",
    ]

    ordering = parsed["请根据时间的先后，一一对应排序"]
    assert ordering.question_type.value == "ordering"
    assert ordering.correct_answer == ["元宵节", "端午节", "中秋节", "重阳节"]


def test_parse_rejects_missing_required_headers() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["题目", "题型", "选项A", "选项B", "正确答案", "解析", "分数", "标签"])
    sheet.append(["题", "单选题", "1", "2", "A", None, 1, None])
    buffer = io.BytesIO()
    workbook.save(buffer)
    rows, errors = AdminQuizService()._parse_import_rows_xlsx(buffer.getvalue())
    assert rows == []
    assert errors[0]["error_code"] == "invalid_workbook_header"


def test_parse_row_level_errors_do_not_abort_batch() -> None:
    rows, errors = _parse([
        ["好题", "单选题", "酸", "甜", None, None, None, None, None, None, "A", None, 8, "分类", None],
        ["坏题型", "选择题", "酸", "甜", None, None, None, None, None, None, "A", None, 8, "分类", None],
        ["坏答案", "单选题", "酸", "甜", None, None, None, None, None, None, "Z", None, 8, "分类", None],
        ["坏配对", "配对题", "左", None, None, None, None, None, None, None, "A", None, 8, "分类", None],
        ["超分数", "单选题", "酸", "甜", "苦", None, None, None, None, None, "A", None, 120, "分类", None],
    ])
    assert len(rows) == 1
    assert len(errors) == 4
    assert {error["row"] for error in errors} == {3, 4, 5, 6}


def test_real_client_template_file_parses() -> None:
    from pathlib import Path

    path = Path(__file__).parent.parent / "fixtures" / "quiz-import-template.xlsx"
    rows, errors = AdminQuizService()._parse_import_rows_xlsx(path.read_bytes())
    # 甲方模板第一个工作表共 11 行示例（第二个 sheet「模板说明」不参与解析）。
    assert errors == []
    assert len(rows) == 11
